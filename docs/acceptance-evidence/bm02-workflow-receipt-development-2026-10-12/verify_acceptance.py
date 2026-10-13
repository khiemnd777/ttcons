"""Inspect canonical Development records and unchanged browser downloads.

Requires the bundled Python runtime and the private temporary API helper.
Writes evidence JSON. Does not edit XLSX, Google Sheets or application records.
Normal API login renewal updates only the private temporary session file.
The known failed material-reuse criterion is reported separately from checks
that establish reproduction and historical content integrity.
"""
import datetime
import hashlib
import importlib.util
import io
import json
from pathlib import Path
import subprocess
import urllib.request
from zipfile import ZipFile
import xml.etree.ElementTree as ET

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from PIL import Image

spec = importlib.util.spec_from_file_location('fixture_helper', '/private/tmp/consflow-bm02-workflow-acceptance.py')
fixture = importlib.util.module_from_spec(spec)
spec.loader.exec_module(fixture)
api, out, root = fixture.api, fixture.OUT, fixture.ROOT
state = json.loads((out / 'fixture-state.json').read_text())
result = {'scope': 'Development: submitted native workflow form, legitimate separate receipt owner and two delivery revisions', 'checked_at': datetime.datetime.now(datetime.timezone.utc).isoformat(), 'checks': {}, 'criteria': {}}


def check(name, condition):
    result['checks'][name] = bool(condition)
    assert condition, name


def read(name):
    return json.loads((out / name).read_text())


def renew(actor):
    private = json.loads(api.PRIVATE.read_text())
    if actor == 'recipient':
        credentials = {key: private[actor][key] for key in ['email', 'password']}
    else:
        values = {}
        for line in (root / 'frontend/.env.development.local').read_text().splitlines():
            if '=' in line and not line.lstrip().startswith('#'):
                key, value = line.split('=', 1)
                values[key.strip()] = value.strip().strip('\"\'')
        credentials = {'email': values['VITE_LOCAL_LOGIN_EMAIL'], 'password': values['VITE_LOCAL_LOGIN_PASSWORD']}
    login = api.call('/auth/login', 'POST', credentials)
    private[actor].update(token=login['token'], user=login['user'])
    api.private(private)


renew('author')
renew('recipient')
before = read('delivery1-before.json')
after = api.call('/work-reports/daily-reports/' + state['source_id'])
form_after = api.call('/work-items/' + state['report_work_item_id'] + '/form')
receipt = api.call('/incoming-acceptances/' + state['receipt_id'])
source = api.call('/projects/' + state['project_id'] + '/daily-reports')
source = next(row for row in source if row['id'] == state['source_id'])
for name, value in [('delivery-after.json', after), ('form-after.json', form_after), ('receipt-after.json', receipt), ('daily-source-after.json', source)]:
    fixture.save(name, value)
r1 = next(row for row in after['revisions'] if row['revision_no'] == 1)
r2 = next(row for row in after['revisions'] if row['revision_no'] == 2)
check('same_workflow_source_and_report_two_revisions', after['work_report_id'] == state['work_report_id'] and after['workflow_form_instance_id'] == state['form_id'] and after['current_revision_no'] == 2 and len(after['revisions']) == 2)
check('revision1_entire_record_unchanged', r1 == before['revisions'][0])
check('revision2_snapshot_distinct', r1['snapshot_sha256'] != r2['snapshot_sha256'])
check('revision1_photo_membership', [p['id'] for p in r1['evidence_records']] == [state['photo1_id']])
check('revision2_photo_membership', [p['id'] for p in r2['evidence_records']] == [state['photo2_id']])
check('recipient_requested_changes', any(e['event_type'] == 'needs_changes' and e['actor_id'] == state['recipient_id'] for e in after['events']))
check('author_submitted_both_revisions', r1['submitted_by'] == r2['submitted_by'] == state['author_id'])
check('form_remains_submitted_revision2', form_after['instance']['status'] == 'submitted' and form_after['instance']['current_revision'] == 2)
old_events = read('form1-submitted.json')['events']
new_events = {e['id']: e for e in form_after['events']}
check('all_revision1_form_events_unchanged', all(new_events.get(e['id']) == e for e in old_events))
check('canonical_form_amendment_event', any(e['event_type'] == 'amended' and e['actor_id'] == state['author_id'] for e in form_after['events']))
keys = {f['field_key'] for f in read('pinned-form-definition.json')['fields']}
check('only_existing_form_answer_keys', set(read('form1-answers.json')) <= keys and set(read('form2-amend-request.json')['answers']) <= keys)
result['snapshot_sha256'] = {'revision1': r1['snapshot_sha256'], 'revision2': r2['snapshot_sha256']}

db_before = read('database-before.json')
db_after = fixture.db_snapshot()
fixture.save('database-after.json', db_after)
check('database_and_migration_state_unchanged', db_before['database'] == db_after['database'] == 'consflow' and db_before['migration_count'] == db_after['migration_count'] == 157)
for table in db_before:
    if table in ['database', 'migration_count']:
        continue
    by_id = {row['id']: row for row in db_after[table]}
    check('all_preexisting_' + table + '_unchanged', all(by_id.get(row['id']) == row for row in db_before[table]))
result['database'] = {'database': db_after['database'], 'migration_count': db_after['migration_count'], 'projects': len(db_after['projects']), 'forms': len(db_after['workflow_form_instances'])}
manifest = read('source-manifest.json')
check('application_source_manifest_unchanged', all(hashlib.sha256((root / path).read_bytes()).hexdigest() == digest for path, digest in manifest.items()))
result['source_file_count'] = len(manifest)
result['head'] = subprocess.check_output(['git', 'rev-parse', 'HEAD'], cwd=root, text=True).strip()
result['containers'] = []
for container in ['ttcons-backend-1', 'admin']:
    record = json.loads(subprocess.check_output(['docker', 'inspect', '--format', '{{json .}}', container], text=True))
    result['containers'].append({'name': container, 'image': record['Image'], 'created': record['Created'], 'started_at': record['State']['StartedAt']})
result['health'] = {path: api.call('/' + path) for path in ['health', 'ready']}

books, archives = {}, {}
result['files'] = {}
for download in read('browser-downloads.json'):
    phase, name = download['phase'], download['file']
    path, raw = out / name, (out / name).read_bytes()
    check(phase + '_copied_original_byte_equal', raw == Path(download['original']).read_bytes())
    book = load_workbook(path)
    books[phase] = book
    with ZipFile(path) as archive:
        archives[phase] = {n: archive.read(n) for n in archive.namelist()}
    check(phase + '_two_expected_sheets', book.sheetnames == ['2026-10-12_01', 'Hình ảnh'])
    check(phase + '_no_excel_error_cells', not [(s.title, c.coordinate) for s in book for row in s for c in row if c.data_type == 'e'])
    photos = []
    for part, data in archives[phase].items():
        if part.startswith('xl/media/'):
            image = Image.open(io.BytesIO(data)).convert('RGB')
            photos.append({'part': part, 'bytes': len(data), 'sha256': hashlib.sha256(data).hexdigest(), 'width': image.width, 'height': image.height, 'pixel_sha256': hashlib.sha256(image.tobytes()).hexdigest()})
    check(phase + '_one_decodable_embedded_photo', len(photos) == 1)
    revision = 2 if phase in ['form2', 'delivery2'] else 1
    expected_photo = Image.open(out / f'BM02-workflow-Lan-{revision}.png').convert('RGB')
    check(phase + '_correct_original_photo_pixels', photos[0]['pixel_sha256'] == hashlib.sha256(expected_photo.tobytes()).hexdigest())
    sheet = book['2026-10-12_01']
    check(phase + '_quantity_notes', f'Khối lượng đã ghi nhận: {9 if revision == 2 else 2}' in sheet['D14'].value)
    check(phase + '_crew', sheet['C25'].value == (6 if revision == 2 else 3))
    check(phase + '_missing_progress_is_not_zero', sheet['D10'].value == 'Chưa có ghi nhận tiến độ trong nguồn báo cáo.')
    check(phase + '_no_invented_assessment', all(sheet[f'C{row}'].value == 'Chưa đánh giá' for row in range(10, 22)))
    result['files'][phase] = {'file': name, 'sha256': hashlib.sha256(raw).hexdigest(), 'bytes': len(raw), 'sheets': book.sheetnames, 'media': photos, 'quantity_note': sheet['D14'].value, 'crew': sheet['C25'].value, 'material_note': sheet['D12'].value}

old_before, old_after = archives['delivery1-before'], archives['delivery1-after']
different_parts = [part for part in sorted(set(old_before) | set(old_after)) if old_before.get(part) != old_after.get(part)]
check('revision1_all_parts_equal_except_mime_declaration_order', set(different_parts) <= {'[Content_Types].xml'})


def content_types(raw):
    declarations = sorted((el.tag, tuple(sorted(el.attrib.items())), el.text or '', el.tail or '') for el in ET.fromstring(raw))
    assert len(declarations) == len(set(declarations)), 'Duplicate MIME declaration'
    return declarations


check('revision1_mime_declarations_equal', content_types(old_before['[Content_Types].xml']) == content_types(old_after['[Content_Types].xml']))
result['revision1_archive_differences'] = different_parts
for revision in [1, 2]:
    form_sheet = books[f'form{revision}']['2026-10-12_01']
    delivery_sheet = books['delivery2' if revision == 2 else 'delivery1-after']['2026-10-12_01']
    check(f'revision{revision}_form_delivery_original_template_cells_match', all(form_sheet.cell(row, col).value == delivery_sheet.cell(row, col).value for row in range(1, 29) for col in range(1, 9)))

result['protected_photos'] = []
recipient_token = json.loads(api.PRIVATE.read_text())['recipient']['token']
for revision, evidence_id in [(1, state['photo1_id']), (2, state['photo2_id'])]:
    request = urllib.request.Request(api.BASE + f'/work-reports/{state["work_report_id"]}/revisions/{revision}/evidence/{evidence_id}/file', headers={'Authorization': 'Bearer ' + recipient_token})
    with urllib.request.urlopen(request, timeout=30) as response:
        raw = response.read()
    check(f'recipient_protected_photo{revision}_original_bytes', raw == (out / f'BM02-workflow-Lan-{revision}.png').read_bytes())
    result['protected_photos'].append({'revision': revision, 'evidence_id': evidence_id, 'sha256': hashlib.sha256(raw).hexdigest()})

result['google'] = []
for phase in ['form1', 'form2', 'delivery2', 'delivery1-after']:
    google = read('google-' + phase + '.json')['cells']
    observed, errors = {}, []
    for s in google['sheets']:
        title = s['properties']['title']
        observed[title] = {}
        for block in s.get('data', []):
            for row_no, row in enumerate(block.get('rowData', []), block.get('startRow', 0) + 1):
                for col_no, cell in enumerate(row.get('values', []), block.get('startColumn', 0) + 1):
                    address = get_column_letter(col_no) + str(row_no)
                    observed[title][address] = cell
                    if cell.get('effectiveValue', {}).get('errorValue'):
                        errors.append((title, address, cell['effectiveValue']))
    check(phase + '_google_no_effective_error_values', not errors)
    differences, populated = [], 0
    for sheet in books[phase]:
        last_row = 75 if sheet.title == '2026-10-12_01' else 5
        for row in range(1, last_row + 1):
            for col in range(1, 9):
                cell = sheet.cell(row, col)
                expected = cell.value if cell.value not in [None, ''] else None
                entered = observed[sheet.title].get(cell.coordinate, {}).get('userEnteredValue', {})
                actual = entered.get('formulaValue', entered.get('numberValue', entered.get('stringValue')))
                actual = actual if actual not in [None, ''] else None
                if expected is not None:
                    populated += 1
                if actual != expected:
                    differences.append((sheet.title, cell.coordinate, expected, actual))
    check(phase + '_google_all_bounded_cells_formulas_match_original', not differences)
    crew = observed['2026-10-12_01']['C28'].get('effectiveValue', {}).get('numberValue')
    check(phase + '_google_crew_formula_recalculated', crew == (6 if phase in ['form2', 'delivery2'] else 3))
    result['google'].append({'phase': phase, 'url': google['spreadsheetUrl'], 'ranges': ["'2026-10-12_01'!A1:H75", "'Hình ảnh'!A1:H5"], 'populated_cells_compared': populated, 'effective_error_count': len(errors), 'crew_sum': crew})

check('valid_receipt_same_project_and_date_submitted', receipt['status'] == 'submitted' and receipt['project_id'] == state['project_id'] and receipt['received_at'].startswith('2026-10-12'))
check('valid_receipt_delivered12_accepted12', float(receipt['items'][0]['delivered_quantity']) == float(receipt['items'][0]['accepted_quantity']) == 12)
check('canonical_distinct_receiving_reporting_owners', receipt['work_item_id'] == state['receipt_work_item_id'] != state['report_work_item_id'] == form_after['instance']['work_item_id'])
check('no_manual_arrival_rows_entered', not read('form1-answers.json').get('material_arrival_items') and not read('form2-amend-request.json')['answers'].get('material_arrival_items'))
check('failed_receipt_reuse_reproduced_in_both_frozen_revisions', all(not row['snapshot_json']['structured_json'].get('material_arrival_items') for row in [r1, r2]))
check('failed_receipt_reuse_reproduced_in_all_five_actual_downloads', all(book['2026-10-12_01']['D12'].value == 'Chưa có ghi nhận vật tư nhập trường trong nguồn báo cáo.' and all(book['2026-10-12_01'].cell(row, col).value in [None, ''] for row in range(25, 28) for col in range(5, 8)) for book in books.values()))
result['criteria'] = {'submitted_workflow_content_and_photo': 'PASS', 'old_delivery_content_photo_immutable_after_form_amendment': 'PASS', 'google_original_file_cells_formulas_images': 'PASS', 'automatic_legitimate_material_receipt_reuse': 'FAIL', 'user_visual_acceptance': 'PENDING'}
result['status'] = 'BLOCKED: automatic material receipt reuse excludes canonical separate owner'
result['screenshots'] = {path.name: hashlib.sha256(path.read_bytes()).hexdigest() for path in sorted(out.glob('*.jpg'))}
fixture.save('acceptance-checks.json', result)
print(json.dumps({'status': result['status'], 'checks_passed': len(result['checks']), 'criteria': result['criteria'], 'database': result['database'], 'google': result['google']}, ensure_ascii=False, indent=2))
