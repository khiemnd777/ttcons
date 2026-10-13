"""Approved Development BM03 client UAT. Domain writes must use canonical APIs.

SQL here is read-only evidence; credentials remain in the established private
helper. Browser exports must be copied unchanged, never regenerated or repaired.
"""
import datetime
import hashlib
import importlib.util
import json
from pathlib import Path
import subprocess
import sys
import urllib.request
import uuid

ROOT = Path(__file__).resolve().parents[3]
OUT = Path(__file__).resolve().parent
spec = importlib.util.spec_from_file_location('private_api', '/private/tmp/consflow-bm02-receipt-mapping-uat.py')
private = importlib.util.module_from_spec(spec)
spec.loader.exec_module(private)
api = private.api


def save(name, value):
    (OUT / name).write_text(json.dumps(value, ensure_ascii=False, indent=2) + '\n')


def state():
    return json.loads((OUT / 'fixture-state.json').read_text())


def persist(value):
    save('fixture-state.json', value)


def upload_file(path, file, fields):
    boundary = 'bm03-uat-' + uuid.uuid4().hex
    body = b''
    for key, value in fields.items():
        body += f'--{boundary}\r\nContent-Disposition: form-data; name="{key}"\r\n\r\n{value}\r\n'.encode()
    mime = 'application/pdf' if file.suffix == '.pdf' else 'image/png'
    body += f'--{boundary}\r\nContent-Disposition: form-data; name="file"; filename="{file.name}"\r\nContent-Type: {mime}\r\n\r\n'.encode() + file.read_bytes() + f'\r\n--{boundary}--\r\n'.encode()
    token = json.loads(api.PRIVATE.read_text())['author']['token']
    request = urllib.request.Request(api.BASE + path, data=body, method='POST', headers={'Authorization': 'Bearer ' + token, 'Content-Type': 'multipart/form-data; boundary=' + boundary})
    with urllib.request.urlopen(request, timeout=30) as response:
        return json.load(response)['data']


def download_protected(path, file):
    token = json.loads(api.PRIVATE.read_text())['author']['token']
    request = urllib.request.Request(api.BASE + path, headers={'Authorization': 'Bearer ' + token})
    with urllib.request.urlopen(request, timeout=30) as response:
        file.write_bytes(response.read())


def make_test_documents():
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib.styles import getSampleStyleSheet
    styles = getSampleStyleSheet()
    rows = [['Item', 'Unit', 'Quantity', 'Price (VND)'], ['Waterproof floor - internal test area', 'm2', '25', '100,000']]
    for filename, title in [('synthetic-client-boq.pdf', 'Development test BOQ source'), ('synthetic-returned-contract.pdf', 'Development test returned document')]:
        file = OUT / filename
        if file.exists():
            continue
        story = [Paragraph(title, styles['Title']), Spacer(1, 20), Paragraph('SYNTHETIC DEVELOPMENT DATA - NOT A REAL CLIENT DOCUMENT', styles['Heading2']), Spacer(1, 12), Paragraph('Project: BM03-REV10-20261013. This source is for internal workflow validation only. No actual contract, delivery, client signature or payment is represented.', styles['BodyText']), Spacer(1, 20)]
        table = Table(rows, colWidths=[250, 45, 65, 100])
        table.setStyle(TableStyle([('BACKGROUND', (0,0), (-1,0), colors.HexColor('#24466c')), ('TEXTCOLOR', (0,0), (-1,0), colors.white), ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'), ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#b8c5d3')), ('TOPPADDING', (0,0), (-1,-1), 10), ('BOTTOMPADDING', (0,0), (-1,-1), 10)]))
        story += [table, Spacer(1, 24), Paragraph('Internal test record / 13 October 2026', styles['BodyText'])]
        SimpleDocTemplate(str(file), pagesize=A4, leftMargin=40, rightMargin=40, topMargin=40, bottomMargin=40).build(story)
    print('Created explicit synthetic test source PDFs')


def form_for(work_item):
    return api.call('/work-items/' + work_item + '/form')


def ensure_document(work_item, label, filename='synthetic-client-boq.pdf'):
    st = state()
    documents = st.setdefault('evidence_documents', {})
    if not documents.get(work_item):
        r = upload_file('/work-items/' + work_item + '/evidence/upload', OUT / filename, {'evidence_type': 'document', 'name': '[Synthetic] ' + label, 'caption': '[Synthetic] Internal Development test source, not a client record.'})
        documents[work_item] = r['id']
        persist(st)
        save(label + '-evidence-document.json', r)
    return documents[work_item]


def submit_form(work_item, answers, label):
    f = form_for(work_item)
    if f['instance']['status'] != 'submitted':
        fid = f['instance']['id']
        r = api.call('/workflow-form-instances/' + fid + '/draft', 'PUT', {'answers': answers, 'expected_answers': f.get('answers', {}), 'comment': '[Synthetic] Canonical Development source capture'})
        save(label + '-draft.json', r)
        current = form_for(work_item)
        r = api.call('/workflow-form-instances/' + fid + '/submit', 'POST', {'answers': answers, 'expected_answers': current.get('answers', {}), 'comment': '[Synthetic] Chốt nguồn thử Development'})
        save(label + '-submitted.json', r)
    return form_for(work_item)


def complete_work(work_item, label):
    w = api.call('/work-items/' + work_item)
    if w['status'] not in ['completed', 'done']:
        if w['status'] in ['pending', 'ready', 'needs_changes']:
            api.call('/work-items/' + work_item + '/start', 'POST', {})
        r = api.call('/work-items/' + work_item + '/complete', 'POST', {'comment': '[Synthetic] Complete canonical test source; no waiver'})
        save(label + '-completed.json', r)
    return api.call('/work-items/' + work_item)


def intake():
    private.renew()
    st = state()
    work_item = '9d285b24-5b69-4e03-b039-398a207471f0'
    if not st.get('intake_boq_evidence'):
        r = upload_file('/work-items/' + work_item + '/evidence/upload', OUT / 'synthetic-client-boq.pdf', {'evidence_type': 'document', 'name': '[Synthetic] BOQ đầu vào thử Development', 'caption': '[Synthetic] Test source; not an actual client document'})
        st['intake_boq_evidence'] = r['id']
        persist(st)
        save('intake-boq-uploaded.json', r)
    submit_form(work_item, {'entry_route': 'client_boq', 'client_boq_evidence': st['intake_boq_evidence']}, 'intake')
    complete_work(work_item, 'intake')
    r = api.call('/projects/' + st['project_id'] + '/pipeline-run')
    save('run-after-intake.json', r)
    print(json.dumps({'current_node': r.get('current_node_id'), 'nodes': [{k:n.get(k) for k in ['id','business_code','status']} for n in r['nodes']]}, ensure_ascii=False))


def advance_current(expected_code):
    st = state()
    run = api.call('/projects/' + st['project_id'] + '/pipeline-run')
    node = next(n for n in run['nodes'] if n['id'] == run['current_node_id'])
    if node['business_code'] == expected_code:
        r = api.call('/pipeline-run-nodes/' + node['id'] + '/advance', 'POST', {'comment': '[Synthetic] Canonical progression; no bypass'})
        save('advance-' + expected_code + '.json', r)
    run = api.call('/projects/' + st['project_id'] + '/pipeline-run')
    save('run-after-' + expected_code + '.json', run)
    node = next(n for n in run['nodes'] if n['id'] == run['current_node_id'])
    print('Current stage:', node['business_code'])


def boq():
    private.renew()
    st = state()
    project = st['project_id']
    for kind in ['supplier', 'client']:
        key = kind + '_partner_id'
        if not st.get(key):
            r = api.call('/partners', 'POST', {'partner_code': 'BM03-REV10-' + kind.upper() + '-20261013', 'name': '[Synthetic] ' + kind + ' Development BM03', 'partner_types': [kind], 'notes': '[Synthetic] Internal test counterparty; no real commercial party.'})
            st[key] = r['id']
            persist(st)
            save(kind + '-partner.json', r)
    if not st.get('boq_id'):
        r = api.call('/projects/' + project + '/boq', 'POST', {'category': '[Synthetic] Flooring test', 'work_item_name': '[Synthetic] Waterproof test floor', 'description': '[Synthetic] Internal source consistent with uploaded test BOQ', 'unit': 'm2', 'quantity': 25, 'expected_unit_price': 100000})
        st['boq_id'] = r['id']
        persist(st)
        save('boq-created.json', r)
    item = api.call('/boq-items/' + st['boq_id'])
    if item['status'] == 'draft':
        r = api.call('/projects/' + project + '/boq/submit', 'POST', {})
        save('boq-domain-submitted.json', r)
    item = api.call('/boq-items/' + st['boq_id'])
    if item['status'] == 'submitted':
        if not item.get('norm_price_locked'):
            r = api.call('/boq-items/' + st['boq_id'] + '/unit-price/lock-norm', 'POST', {'unit_price': 100000, 'comment': '[Synthetic] Test norm from source BOQ'})
            save('boq-norm-locked.json', r)
        r = api.call('/projects/' + project + '/boq/approve', 'POST', {})
        save('boq-domain-approved.json', r)
    if not st.get('quotation_id'):
        r = api.call('/projects/' + project + '/supplier-quotations', 'POST', {'supplier_partner_id': st['supplier_partner_id'], 'quotation_code': 'BM03-REV10-Q-20261013', 'quotation_date': '2026-10-13', 'valid_until': '2026-11-13', 'currency': 'VND', 'items': [{'boq_item_id': st['boq_id'], 'quantity': 25, 'unit_price': 90000, 'compliance_status': 'compliant'}], 'notes': '[Synthetic] Internal Development price comparison source; no supplier contact.'})
        st['quotation_id'] = r['id']
        persist(st)
        save('quotation-created.json', r)
    quote = api.call('/supplier-quotations/' + st['quotation_id'])
    if quote['status'] == 'draft':
        save('quotation-submitted.json', api.call('/supplier-quotations/' + st['quotation_id'] + '/submit', 'POST', {}))
    quote = api.call('/supplier-quotations/' + st['quotation_id'])
    if quote['status'] == 'submitted':
        save('quotation-qualified.json', api.call('/supplier-quotations/' + st['quotation_id'] + '/qualify', 'POST', {'technical_status': 'compliant', 'commercial_status': 'compliant', 'comment': '[Synthetic] Price and quantity match the test scenario.'}))
    quote = api.call('/supplier-quotations/' + st['quotation_id'])
    item = api.call('/boq-items/' + st['boq_id'])
    if not item.get('final_price_locked'):
        qi = quote['items'][0]
        save('quotation-selected.json', api.call('/supplier-quotation-items/' + qi['id'] + '/select', 'POST', {'decision_reason': '[Synthetic] Qualified test source; canonical selection, no waiver.'}))
    work_item = '4c0160fd-336e-4f8a-8dd0-82f265805b41'
    submit_form(work_item, {'sent_to': '[Synthetic] Internal test recipient', 'sent_channel': 'meeting'}, 'boq-form')
    ensure_document(work_item, 'boq-source')
    complete_work(work_item, 'boq-work')
    advance_current('boq_without_design')
    save('procurement-generated.json', api.call('/procurement-requests?project_id=' + project))


def prepare_contract(kind):
    st = state()
    project = st['project_id']
    key = kind + '_contract_id'
    partner = st['client_partner_id'] if kind == 'outbound' else st['supplier_partner_id']
    if not st.get(key):
        r = api.call('/projects/' + project + '/contracts', 'POST', {'contract_code': 'BM03-REV10-' + kind.upper(), 'title': '[Synthetic] ' + kind + ' internal test contract', 'contract_type': 'main_client' if kind == 'outbound' else 'supplier', 'direction': 'receivable' if kind == 'outbound' else 'payable', 'counterparty_partner_id': partner, 'currency': 'VND', 'original_amount': 2250000, 'effective_date': '2026-10-13'})
        st[key] = r['id']
        persist(st)
        save(kind + '-contract-created.json', r)
    version_key = kind + '_contract_version_id'
    if not st.get(version_key):
        r = api.call('/contracts/' + st[key] + '/versions', 'POST', {'version_label': 'UAT V1', 'amount': 2250000, 'scope_summary': '[Synthetic] Internal test floor25m2; no legal contract or actual payment.', 'commercial_terms': {}, 'payment_terms': {}, 'comment': '[Synthetic] Canonical test version'})
        st[version_key] = next(v['id'] for v in r['versions'] if v['version_label'] == 'UAT V1')
        persist(st)
        save(kind + '-contract-version-created.json', r)
    c = api.call('/contracts/' + st[key])
    if st[version_key] == st[key]:
        st[version_key] = next(v['id'] for v in c['versions'] if v['version_label'] == 'UAT V1')
        persist(st)
    v = next(v for v in c['versions'] if v['id'] == st[version_key])
    if v['status'] == 'draft':
        save(kind + '-contract-version-submitted.json', api.call('/contract-versions/' + v['id'] + '/submit', 'POST', {'comment': '[Synthetic] Internal review'}))
    c = api.call('/contracts/' + st[key])
    v = next(v for v in c['versions'] if v['id'] == st[version_key])
    if v['status'] in ['submitted', 'review']:
        save(kind + '-contract-version-approved.json', api.call('/contract-versions/' + v['id'] + '/approve', 'POST', {'comment': '[Synthetic] Internal test approval; no real legal agreement.'}))
    c = api.call('/contracts/' + st[key])
    save(kind + '-contract-prepared.json', c)
    print(json.dumps({'contract': st[key], 'versions': [{k:v.get(k) for k in ['id','status','approval_request_id']} for v in c['versions']]}, ensure_ascii=False))
    return c


def outbound_prepare():
    private.renew()
    prepare_contract('outbound')
    work_item = '37d90fa3-abe5-4bbe-af17-34b1990730fb'
    w = api.call('/work-items/' + work_item)
    if w['status'] == 'ready':
        api.call('/work-items/' + work_item + '/start', 'POST', {})
    ensure_document(work_item, 'outbound-test-document', 'synthetic-returned-contract.pdf')
    f = form_for(work_item)
    if f['instance']['status'] != 'submitted':
        answers = {'hardcopy_set': '[Synthetic] Internal Development test contract packet; no real client dispatch.', 'receiver': '[Synthetic] Internal UAT recipient', 'delivery_method': 'hand_delivery'}
        r = api.call('/workflow-form-instances/' + f['instance']['id'] + '/draft', 'PUT', {'answers': answers, 'expected_answers': f.get('answers', {})})
        save('outbound-form-draft-prepared.json', r)
    save('outbound-before-browser.json', api.call('/contracts/' + state()['outbound_contract_id']))


def outbound_finish():
    private.renew()
    st = state()
    c = api.call('/contracts/' + st['outbound_contract_id'])
    assert c['status'] in ['signed', 'active'] and c['current_version_id'] == st['outbound_contract_version_id']
    exports = json.loads((OUT / 'outbound-native-pdf-generated.json').read_text())['pdf_exports']
    export = exports[0]
    file = OUT / 'outbound-native-signed-uat.pdf'
    if not file.exists():
        download_protected('/workflow-form-pdf-exports/' + export['id'] + '/file', file)
    assert hashlib.sha256(file.read_bytes()).hexdigest() == export['checksum_sha256']
    if not st.get('outbound_dispatch_id'):
        r = api.call('/contracts/' + st['outbound_contract_id'] + '/dispatches', 'POST', {'contract_version_id': st['outbound_contract_version_id'], 'recipient_name': '[Synthetic] Internal UAT recipient', 'delivery_method': 'other', 'note': '[Synthetic] Internal test dispatch only; no actual delivery.'})
        dispatch = r['dispatches'][0] if 'dispatches' in r else r
        st['outbound_dispatch_id'] = dispatch['id']
        persist(st)
        save('outbound-dispatch-created.json', r)
    if not st.get('outbound_scan_document_id'):
        r = upload_file('/contract-dispatches/' + st['outbound_dispatch_id'] + '/documents/upload', file, {'name': '[Synthetic] Returned internal UAT signed PDF'})
        st['outbound_scan_document_id'] = r['id']
        persist(st)
        save('outbound-scan-uploaded.json', r)
    c = api.call('/contracts/' + st['outbound_contract_id'])
    dispatch = next(d for d in c['dispatches'] if d['id'] == st['outbound_dispatch_id'])
    if dispatch['status'] != 'received':
        save('outbound-scan-received.json', api.call('/contract-dispatches/' + st['outbound_dispatch_id'] + '/receive', 'POST', {'returned_scan_document_id': st['outbound_scan_document_id'], 'note': '[Synthetic] Internal receipt of actual system-generated UAT signature PDF; no legal client signature.'}))
    c = api.call('/contracts/' + st['outbound_contract_id'])
    if c['status'] != 'active':
        save('outbound-activated.json', api.call('/contracts/' + st['outbound_contract_id'] + '/activate', 'POST', {'comment': '[Synthetic] Canonical internal UAT activation'}))
    complete_work('37d90fa3-abe5-4bbe-af17-34b1990730fb', 'outbound-work')
    advance_current('outbound_contract')


def input_chain():
    private.renew()
    st = state()
    complete_work('be5da168-47b4-4e7d-8667-b2eea0aeae26', 'input-takeoff')
    complete_work('4325f888-24cf-4c0b-882a-5b3926cb703d', 'input-sourcing')
    advance_current('input_takeoff_purchase')
    complete_work('fb9892ba-3fbc-461f-af2e-a1cdf154ddde', 'input-negotiation')
    advance_current('price_negotiation')
    requests = api.call('/procurement-requests?project_id=' + st['project_id'])
    request = requests[0]
    st['procurement_request_id'] = request['id']
    persist(st)
    if request['status'] == 'draft':
        save('procurement-submitted.json', api.call('/procurement-requests/' + request['id'] + '/submit', 'POST', {}))
    request = api.call('/procurement-requests/' + request['id'])
    if request['status'] in ['submitted', 'waiting_approval']:
        save('procurement-approved.json', api.call('/procurement-requests/' + request['id'] + '/approve', 'POST', {}))
    prepare_contract('input')
    complete_work('52abeed6-a354-4e1e-8b63-d51ee306b248', 'input-contract')
    advance_current('input_contract_advance')


def make_inspection_assets():
    from PIL import Image, ImageDraw
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet
    image_file = OUT / 'synthetic-inspection.png'
    if not image_file.exists():
        image = Image.new('RGB', (960, 640), '#f2f5f9')
        draw = ImageDraw.Draw(image)
        for box, color in [((0,0,80,80),'#ec3131'),((880,0,959,80),'#20bc44'),((0,560,80,639),'#235ddd'),((880,560,959,639),'#ffd326')]:
            draw.rectangle(box, fill=color)
        for x in range(80, 900, 80):
            draw.line((x,80,x,560), fill='#c3cfdb', width=2)
        for y in range(80, 561, 80):
            draw.line((80,y,880,y), fill='#c3cfdb', width=2)
        draw.text((300,280), 'SYNTHETIC / DEVELOPMENT UAT', fill='#1a3b5b')
        draw.text((330,320), 'FOUR CORNER IMAGE CHECK', fill='#1a3b5b')
        image.save(image_file)
    file = OUT / 'synthetic-qc-plan.pdf'
    if not file.exists():
        styles = getSampleStyleSheet()
        story = [Paragraph('Development inspection plan', styles['Title']), Spacer(1,20), Paragraph('SYNTHETIC INTERNAL TEST - NOT A REAL CONSTRUCTION RECORD', styles['Heading2']), Spacer(1,16), Paragraph('Project BM03-REV10-20261013. Supervisor checks the actual test image: all four colored corners and the SYNTHETIC label must be visible. Pass means this image criterion passed; it is not a statement of actual site quality.', styles['BodyText']), Spacer(1,16), Paragraph('Date: 13 October 2026. Evidence: synthetic-inspection.png. Responsibility: internal UAT supervisor.', styles['BodyText'])]
        SimpleDocTemplate(str(file), pagesize=A4, leftMargin=40, rightMargin=40, topMargin=40, bottomMargin=40).build(story)
    print('Prepared internal image-inspection scenario')


def planning():
    private.renew()
    st = state()
    plan_work = '34282293-b45a-477c-b793-7c0035cc5aab'
    handoff_work = '1d6e55cc-ee66-4f62-9d02-6f51b5cf60c3'
    evidence = ensure_document(plan_work, 'planning-qc-source', 'synthetic-qc-plan.pdf')
    answers = {'meeting_title': '[Synthetic] Kế hoạch kiểm tra ảnh thử nội bộ', 'meeting_date': '2026-10-13', 'chair': '[Synthetic] Điều phối UAT', 'attendees': [{'name': '[Synthetic] Giám sát nội bộ', 'role': 'Giám sát', 'attendance': 'present'}], 'agenda': '[Synthetic] Thực hiện kiểm tra hiển thị bốn mốc màu của ảnh thử; không phải kiểm tra chất lượng công trình thật.', 'qa_items': [{'item': '[Synthetic] Đã chuẩn bị ảnh và tiêu chí kiểm tra bốn góc', 'status': 'done', 'evidence_id': evidence, 'note': '[Synthetic] Nguồn thử nội bộ đã được tạo và đọc.'}]}
    submit_form(plan_work, answers, 'planning')
    ensure_document(handoff_work, 'handoff-qc-source', 'synthetic-qc-plan.pdf')
    submit_form(handoff_work, {'handoff_title': '[Synthetic] Đợt kiểm tra ảnh thử', 'inspection_items': [{'title': '[Synthetic] Kiểm tra hiển thị ảnh bốn góc', 'location': '[Synthetic] Vùng thử Development', 'criteria': '[Synthetic] Ảnh thử phải hiển thị đủ bốn mốc màu góc và nhãn SYNTHETIC; không phải ảnh công trình.', 'scheduled_date': '2026-10-13', 'responsible_role': 'Giám sát'}], 'handoff_notes': '[Synthetic] Chỉ nghiệm thu kịch bản thử nội bộ.'}, 'handoff')
    r = api.call('/projects/' + st['project_id'] + '/construction-handoff')
    save('handoff-candidates.json', r)
    print(json.dumps(r, ensure_ascii=False)[:9000])


def activate_planning():
    private.renew()
    st = state()
    preview = api.call('/projects/' + st['project_id'] + '/construction-handoff')
    save('handoff-preview-final.json', preview)
    if not preview['handoffs']:
        assert preview['can_activate'] and preview['source_form_instance_id']
        actor = json.loads(api.PRIVATE.read_text())['author']['user']['id']
        assert any(p['id'] == actor for p in preview['participants'])
        items = [{**item, 'assigned_user_id': actor} for item in preview['candidates']]
        assert items and all(item['evidence_required'] for item in items)
        request = {'title': preview['suggested_title'], 'notes': '[Synthetic] Activate exact submitted source; retain mandatory evidence.', 'source_run_node_id': preview['source_run_node_id'], 'source_form_instance_id': preview['source_form_instance_id'], 'source_form_revision_no': preview['source_form_revision_no'], 'items': items}
        save('handoff-activation-request.json', request)
        save('handoff-activated.json', api.call('/projects/' + st['project_id'] + '/construction-handoff/activate', 'POST', request))
    current = api.call('/projects/' + st['project_id'] + '/construction-handoff')
    save('handoff-activation-readback.json', current)
    assert len(current['handoffs']) == 1
    st['handoff_id'] = current['handoffs'][0]['id']
    st['generated_check_id'] = current['handoffs'][0]['items'][0]['generated_check_id']
    st['generated_check_work_item_id'] = current['handoffs'][0]['items'][0]['generated_work_item_id']
    persist(st)
    complete_work('34282293-b45a-477c-b793-7c0035cc5aab', 'planning')
    complete_work('1d6e55cc-ee66-4f62-9d02-6f51b5cf60c3', 'handoff')
    advance_current('construction_planning')


def prepare_execution_photo():
    private.renew()
    st = state()
    wid = st['generated_check_work_item_id']
    work = api.call('/work-items/' + wid)
    assert work['status'] in ['ready', 'in_progress']
    if not st.get('inspection_photo_id'):
        photo = upload_file('/work-items/' + wid + '/evidence/upload', OUT / 'synthetic-inspection.png', {'evidence_type': 'photo', 'name': '[Synthetic] Ảnh kiểm tra bốn góc', 'caption': '[Synthetic] Test image; not physical site evidence.'})
        st['inspection_photo_id'] = photo['id']
        persist(st)
        save('execution-inspection-photo-uploaded.json', photo)
    workspace = api.call('/projects/' + st['project_id'] + '/construction-workspace')
    check = next(c for c in workspace['checks'] if c['id'] == st['generated_check_id'])
    save('execution-check-before-completion.json', check)
    print(json.dumps({'work_item_id': wid, 'photo_id': st['inspection_photo_id'], 'check_id': st['generated_check_id']}))


def correct_test_schedule():
    private.renew()
    st = state()
    wid = '1d6e55cc-ee66-4f62-9d02-6f51b5cf60c3'
    form = form_for(wid)
    save('handoff-before-schedule-amendment.json', form)
    answers = json.loads(json.dumps(form['answers']))
    # Test execution is immediate. The optional schedule was an unintended fixture assumption.
    for row in answers['inspection_items']:
        row.pop('scheduled_date', None)
    if form['instance']['current_revision'] == 1:
        assert form['capabilities']['can_amend']
        save('handoff-schedule-amended.json', api.call('/workflow-form-instances/' + form['instance']['id'] + '/amend', 'POST', {'answers': answers, 'expected_revision': 1, 'reason': '[Synthetic] Correct test fixture schedule: internal check runs immediately, no 09:00 appointment; retain original revision and evidence gate.'}))
    preview = api.call('/projects/' + st['project_id'] + '/construction-handoff')
    save('handoff-r2-preview.json', preview)
    assert preview['source_form_revision_no'] == 2
    actor = json.loads(api.PRIVATE.read_text())['author']['user']['id']
    items = [{**item, 'assigned_user_id': actor} for item in preview['candidates']]
    assert all(item['scheduled_at'] == '' and item['evidence_required'] for item in items)
    request = {'title': preview['suggested_title'], 'notes': '[Synthetic] Revised immediate internal test schedule, no early-start waiver.', 'source_run_node_id': preview['source_run_node_id'], 'source_form_instance_id': preview['source_form_instance_id'], 'source_form_revision_no': 2, 'items': items}
    activation = api.call('/projects/' + st['project_id'] + '/construction-handoff/activate', 'POST', request)
    save('handoff-r2-activated.json', activation)
    preview = api.call('/projects/' + st['project_id'] + '/construction-handoff')
    save('handoff-r2-readback.json', preview)
    latest = next(h for h in preview['handoffs'] if h['source_form_revision_no'] == 2)
    st.update(handoff_id=latest['id'], generated_check_id=latest['items'][0]['generated_check_id'], generated_check_work_item_id=latest['items'][0]['generated_work_item_id'])
    persist(st)
    save('handoff-after-schedule-amendment.json', form_for(wid))
    print(json.dumps({'revision': 2, 'new_check_work_item': st['generated_check_work_item_id'], 'old_revision_retained': True}))


def finish_execution():
    private.renew()
    st = state()
    workspace = api.call('/projects/' + st['project_id'] + '/construction-workspace')
    check = next(c for c in workspace['checks'] if c['id'] == st['generated_check_id'])
    if check['status'] == 'pending':
        # The root inspected the uploaded synthetic image with the standard viewer.
        result = api.call('/construction-checks/' + st['generated_check_id'] + '/complete', 'POST', {'result': 'pass', 'note': '[Synthetic] Four colored corners and SYNTHETIC label observed in the actual uploaded test image. No claim about real site quality.', 'evidence_ids': [st['inspection_photo_id']]})
        save('execution-inspection-completed.json', result)
    daily = 'd89a0005-e96f-473b-886d-59fb7fb7cfbd'
    w = api.call('/work-items/' + daily)
    if w['status'] == 'ready':
        api.call('/work-items/' + daily + '/start', 'POST', {})
    submit_form(daily, {'report_date': '2026-10-13', 'reported_by': '[Synthetic] Giám sát UAT nội bộ', 'construction_scope': 'flooring', 'weather': 'normal', 'completed_items': [{'title': '[Synthetic] Kiểm tra ảnh thử đủ bốn góc màu và nhãn SYNTHETIC', 'result': 'done'}]}, 'execution-daily')
    complete_work(daily, 'execution-daily')
    save('execution-work-items-final.json', api.call('/projects/' + st['project_id'] + '/work-items'))
    advance_current('construction_execution')


def prepare_client():
    private.renew()
    st = state()
    wid = '778be55c-f1aa-423a-b888-40807c11c2d5'
    work = api.call('/work-items/' + wid)
    assert work['run_node_id'] == api.call('/projects/' + st['project_id'] + '/pipeline-run')['current_node_id']
    if work['status'] == 'ready':
        api.call('/work-items/' + wid + '/start', 'POST', {})
    if not st.get('client_photo_id'):
        photo = upload_file('/work-items/' + wid + '/evidence/upload', OUT / 'synthetic-inspection.png', {'evidence_type': 'photo', 'name': '[Synthetic] BM03 ảnh nguồn nghiệm thu', 'caption': '[Synthetic] Internal BM03 image test, not physical construction evidence.'})
        st['client_photo_id'] = photo['id']
        persist(st)
        save('client-photo-uploaded.json', photo)
    answers = json.loads((ROOT / 'docs/acceptance-evidence/bm0308-valid-flow-development-2026-10-13/client-form-answers.json').read_text())
    answers['reference_no'] = 'UAT-BM03-REV10-20261013-R1'
    answers['overall_comment'] = '[Synthetic] Rev10 R1; dữ liệu thử nội bộ, không phải nghiệm thu công trình thật.'
    answers['checklist_items'][0]['evidence_id'] = st['client_photo_id']
    form = form_for(wid)
    assert form['instance']['status'] == 'draft'
    save('client-r1-answers.json', answers)
    save('client-r1-draft.json', api.call('/workflow-form-instances/' + form['instance']['id'] + '/draft', 'PUT', {'answers': answers, 'expected_answers': form['answers'], 'comment': '[Synthetic] Prepare R1 for actual UI Chốt hồ sơ.'}))
    save('client-r1-before-ui-submit.json', form_for(wid))
    print('Prepared actual pinned client R1 draft; browser submit required.')


def client_snapshot(label):
    private.renew()
    wid = '778be55c-f1aa-423a-b888-40807c11c2d5'
    for suffix, path in [('work-item', '/work-items/' + wid), ('form', '/work-items/' + wid + '/form'), ('ledger', '/work-items/' + wid + '/subcontractor-reconciliation'), ('history-r1', '/work-items/' + wid + '/subcontractor-reconciliation?revision=1')]:
        save(label + '-' + suffix + '.json', api.call(path))
    print('Saved client current form, ledger, readiness and R1 history:', label)


def sql(query):
    raw = subprocess.check_output(['docker', 'exec', 'ttcons-postgres-1', 'sh', '-c', 'psql -v ON_ERROR_STOP=1 -U "$POSTGRES_USER" -d "$POSTGRES_DB" -At -c "$1"', 'sh', query], text=True)
    return json.loads(raw)


def baseline(prefix):
    existing = sql("SELECT coalesce(json_agg(table_name),'[]') FROM information_schema.columns WHERE table_schema='public' AND column_name='id'")
    tables = [t for t in existing if t.startswith(('pipeline_', 'workflow_form_', 'subcontractor_checklist_', 'owner_client_checklist_')) or t in ['projects', 'project_members', 'work_items', 'acceptance_records', 'acceptance_checklist_items', 'handover_records', 'handover_checklist_items', 'documents', 'evidence_records', 'work_reports', 'work_report_revisions', 'work_report_events', 'boq_items', 'stock_movements']]
    parts = ["'" + t + "',(SELECT coalesce(json_agg(json_build_object('id',id,'md5',md5(row_to_json(r)::text)) ORDER BY id),'[]') FROM " + t + ' r)' for t in sorted(tables)]
    result = sql("SELECT json_build_object('database',current_database(),'migration_count',(SELECT count(*) FROM schema_migrations)," + ','.join(parts) + ')')
    save(prefix + '-database.json', result)
    print(json.dumps({'snapshot': prefix, 'tables': len(tables), 'rows': sum(len(result[t]) for t in tables), 'migrations': result['migration_count']}))


def source_manifest(name):
    paths = subprocess.check_output(['rg', '--files', 'backend/internal', 'backend/cmd', 'backend/migrations', 'frontend/src', 'frontend/tests', 'frontend/package.json', 'frontend/package-lock.json', 'Makefile', 'docker-compose.yml'], cwd=ROOT, text=True).splitlines()
    files = {p: hashlib.sha256((ROOT / p).read_bytes()).hexdigest() for p in sorted(paths)}
    save(name, {'head': subprocess.check_output(['git', 'rev-parse', 'HEAD'], cwd=ROOT, text=True).strip(), 'captured_at': datetime.datetime.now(datetime.timezone.utc).isoformat(), 'files': files})
    print(json.dumps({'manifest': name, 'files': len(files)}))


def storage_manifest(name):
    files = {str(p.relative_to(ROOT / 'storage')): hashlib.sha256(p.read_bytes()).hexdigest() for p in sorted((ROOT / 'storage').rglob('*')) if p.is_file()}
    save(name, files)
    print(json.dumps({'storage_manifest': name, 'files': len(files)}))


def scaffold():
    private.renew()
    state_path = OUT / 'fixture-state.json'
    state = json.loads(state_path.read_text()) if state_path.exists() else {}
    template_id = '60e28e04-1183-4d1f-a357-0c2e821442bf'
    template = api.call('/pipeline-templates/' + template_id)
    assert template['status'] == 'published' and template['engine_version'] == 4 and template['version'] == 10
    save('official-rev10-scaffold.json', template)
    if not state.get('project_id'):
        project = api.call('/projects', 'POST', {'name': '[Synthetic] BM03 gói client Rev10 — nghiệm thu 13-10-2026', 'code': 'BM03-REV10-20261013', 'client_name': '[Synthetic] Chủ đầu tư thử Development', 'address': '[Synthetic] Khu thử nội bộ', 'pipeline_template_id': template_id})
        state['project_id'] = project['id']
        save('fixture-state.json', state)
    project_id = state['project_id']
    if not state.get('member_id'):
        actor = json.loads(api.PRIVATE.read_text())['author']['user']['id']
        members = api.call('/projects/' + project_id + '/members')
        staff = next(r for r in members['options']['staff'] if r['user_id'] == actor)
        roles = sorted({'PM'} | {role for node in template['nodes'] for role in node['responsible_role_ids']} | {role for node in template['nodes'] for step in node['steps'] for role in step['required_role_ids']})
        member = api.call('/projects/' + project_id + '/members', 'POST', {'staff_id': staff['id'], 'role_ids': roles, 'project_role_ids': [], 'notes': '[Synthetic] Existing Admin actor on this test project only; no global RBAC change.'})
        state.update(member_id=member['id'], staff_id=staff['id'], roles=roles)
        save('fixture-state.json', state)
    if not state.get('pm_set'):
        api.call('/projects/' + project_id + '/project-manager', 'PUT', {'staff_id': state['staff_id']})
        state['pm_set'] = True
        save('fixture-state.json', state)
    if not state.get('run_id'):
        readiness = api.call('/projects/' + project_id + '/startup-readiness')
        save('startup-readiness.json', readiness)
        assert readiness['ready_to_start'], readiness
        run = api.call('/projects/' + project_id + '/start-pipeline', 'POST', {'template_id': template_id})
        state['run_id'] = run['id']
        save('fixture-state.json', state)
    items = api.call('/projects/' + project_id + '/work-items')
    save('official-work-items-created.json', items)
    print(json.dumps({'project_id': project_id, 'run_id': state['run_id'], 'work_items': len(items)}))


def verify():
    import openpyxl
    before = json.loads((OUT / 'before-database.json').read_text())
    after = json.loads((OUT / 'after-database.json').read_text())
    changes = []
    checked = 0
    for table, rows in before.items():
        if not isinstance(rows, list):
            continue
        current = {r['id']: r['md5'] for r in after[table]}
        for r in rows:
            checked += 1
            if current.get(r['id']) != r['md5']:
                changes.append({'table': table, 'id': r['id']})
    old_files = json.loads((OUT / 'storage-before.json').read_text())
    new_files = json.loads((OUT / 'storage-after.json').read_text())
    changed_files = [p for p, checksum in old_files.items() if new_files.get(p) != checksum]
    xlsx = json.loads((OUT / 'xlsx-inspection.json').read_text())
    data = json.loads((OUT / 'google-cells.json').read_text())['structuredContent']['sheets'][0]['data'][0]
    workbook = openpyxl.load_workbook(OUT / 'BM03-client-closed-after-fix.xlsx', data_only=False)
    sheet = workbook.worksheets[0]
    mismatches = []
    errors = []
    populated = 0
    for i in range(30):
        row = data.get('rowData', [])[i] if i < len(data.get('rowData', [])) else {}
        values = row.get('values', [])
        for j in range(7):
            cell = values[j] if j < len(values) else {}
            effective = cell.get('effectiveValue', {})
            value = effective.get('stringValue', effective.get('numberValue', effective.get('boolValue')))
            expected = sheet.cell(i + 1, j + 1).value
            if expected is not None:
                populated += 1
            if value != expected:
                mismatches.append([i + 1, j + 1, expected, value])
            if 'errorValue' in effective:
                errors.append([i + 1, j + 1, effective['errorValue']])
    result = {'historical_rows_checked': checked, 'historical_row_changes': changes, 'historical_files_checked': len(old_files), 'historical_file_changes': changed_files, 'migration_count': after['migration_count'], 'download_byte_equal': xlsx['byte_equal'], 'source_cells_unchanged': xlsx['unchanged_values'], 'google_cells_checked': 210, 'google_populated_cells': populated, 'google_mismatches': mismatches, 'google_effective_errors': errors, 'scope': 'Retained original baseline and bounded closed-client export/readonly proof; current full Rev10 positive acceptance is checked separately by verify-rev10.'}
    result['bounded_checks_passed'] = not changes and not changed_files and xlsx['byte_equal'] and xlsx['unchanged_values'] and not mismatches and not errors and populated == 60 and after['migration_count'] == 158
    save('evidence-verification.json', result)
    print(json.dumps(result, ensure_ascii=False))
    assert result['bounded_checks_passed']


def verify_rev10():
    import io
    import zipfile
    import openpyxl
    from PIL import Image
    load = lambda name: json.loads((OUT / name).read_text())
    r1 = load('client-r1-submitted-ledger.json')['reconciliation']
    r2 = load('client-final-ledger.json')['reconciliation']
    history = load('client-final-history-r1.json')
    assertions = {
        'r1_ledger_unchanged_after_r2_and_completion': r1 == history['reconciliation'],
        'history_readonly': history['readonly'] and not history['capabilities']['can_write'],
        'distinct_submitted_revision_ledgers': r1['id'] != r2['id'] and [r1['form_revision_no'], r2['form_revision_no']] == [1, 2] and r1['status'] == r2['status'] == 'submitted',
        'same_source_identity': r1['work_item_id'] == r2['work_item_id'] == '778be55c-f1aa-423a-b888-40807c11c2d5' and r1['form_instance_id'] == r2['form_instance_id'] == 'f5dbac7f-86f4-4b16-a0a1-80092e8592db',
        'r1_ready_after_submit': load('client-r1-ready.json')['execution_target']['source_ready'],
        'r2_not_ready_before_own_submit': not load('client-r2-before-ledger-work-item.json')['execution_target']['source_ready'],
        'r2_ready_after_submit': load('client-r2-submitted-work-item.json')['execution_target']['source_ready'],
        'client_completed': load('client-final-work-item.json')['status'] == 'completed' and load('client-final-work-item.json')['completed_at'] is not None,
        'closed_current_readonly': load('client-final-ledger.json')['readonly'] and not load('client-final-ledger.json')['capabilities']['can_write'],
        'commercial_confirmation_independent': float(r1['items'][0]['confirmed_quantity']) == 22 and float(r2['items'][0]['confirmed_quantity']) == 24,
        'physical_snapshot_independent': r1['items'][0]['payload_json']['actual_quantity'] == 24.8 and r2['items'][0]['payload_json']['actual_quantity'] == 24.2,
        'independent_quantity_items_unchanged': load('client-r1-submitted-ledger.json')['form']['answers']['quantity_items'] == load('client-final-form.json')['answers']['quantity_items'],
    }
    built = load('source-manifest-checkbox-built.json')['files']
    assertions['current_source_matches_built'] = all(hashlib.sha256((ROOT / name).read_bytes()).hexdigest() == checksum for name, checksum in built.items())
    files, mismatch, errors, checked = [], [], [], 0
    media = []
    for label, original in [('r1', 'BM-GS-03 (13).xlsx'), ('r2', 'BM-GS-03 (14).xlsx')]:
        file = OUT / ('BM03-Rev10-' + label.upper() + '.xlsx')
        native = Path('/Users/khiemnguyen/Downloads') / original
        assertions[label + '_download_copy_byte_equal'] = file.read_bytes() == native.read_bytes()
        archive = zipfile.ZipFile(file)
        assertions[label + '_zip_valid'] = archive.testzip() is None
        image_names = [name for name in archive.namelist() if name.startswith('xl/media/')]
        assertions[label + '_one_embedded_image'] = len(image_names) == 1
        image_bytes = archive.read(image_names[0])
        media.append(hashlib.sha256(image_bytes).hexdigest())
        image = Image.open(io.BytesIO(image_bytes)).convert('RGB')
        width, height = image.size
        corners = [image.getpixel((x, y)) for x,y in [(0,0),(width-1,0),(0,height-1),(width-1,height-1)]]
        assertions[label + '_image_aspect_and_corners'] = abs(width / height - 1.5) < 0.001 and corners[0][0] > 180 and corners[1][1] > 140 and corners[2][2] > 180 and corners[3][0] > 180 and corners[3][1] > 140
        workbook = openpyxl.load_workbook(file)
        sheet = workbook.worksheets[0]
        assertions[label + '_physical_numeric_unit'] = sheet['E9'].value == 25 and sheet['F9'].value == (24.8 if label == 'r1' else 24.2) and 'm2' in sheet['F9'].number_format
        assertions[label + '_report_number'] = sheet['B5'].value == 'UAT-BM03-REV10-20261013-' + label.upper()
        google = load('google-' + label + '-cells.json')['structuredContent']
        populated = 0
        for google_sheet in google['sheets']:
            title = google_sheet['properties']['title']
            source_sheet = workbook[title]
            data = google_sheet['data'][0]
            rows = data.get('rowData', [])
            columns = 7 if title == workbook.worksheets[0].title else 8
            for i in range(30):
                values = rows[i].get('values', []) if i < len(rows) else []
                for j in range(columns):
                    checked += 1
                    cell = values[j] if j < len(values) else {}
                    effective = cell.get('effectiveValue', {})
                    actual = effective.get('stringValue', effective.get('numberValue', effective.get('boolValue')))
                    expected = source_sheet.cell(i + 1, j + 1).value
                    if expected is not None:
                        populated += 1
                    if actual != expected:
                        mismatch.append([label, title, i + 1, j + 1, expected, actual])
                    if 'errorValue' in effective or str(cell.get('formattedValue', '')).startswith('#VALUE'):
                        errors.append([label, title, i + 1, j + 1, cell])
        files.append({'revision': label, 'download': str(native), 'copy': str(file), 'sha256': hashlib.sha256(file.read_bytes()).hexdigest(), 'media_sha256': media[-1], 'image_dimensions': [width, height], 'populated_cells': populated, 'google_url': google['spreadsheetUrl']})
    assertions['same_image_bytes_r1_r2'] = media[0] == media[1]
    assertions['google_source_values_equal'] = not mismatch
    assertions['google_no_effective_errors'] = not errors
    result = {'assertions': assertions, 'source_files_checked': len(built), 'google_cells_checked': checked, 'google_mismatches': mismatch, 'google_errors': errors, 'files': files, 'scope': 'Synthetic Development full published engine4 Rev10 upstream to own client BM03 completion; Admin test actor. No whole-project/finance/legal/site-quality or user UX acceptance.'}
    result['passed'] = all(assertions.values())
    save('rev10-evidence-verification.json', result)
    print(json.dumps(result, ensure_ascii=False))
    assert result['passed']


if __name__ == '__main__':
    if sys.argv[1] == 'baseline':
        baseline(sys.argv[2])
    elif sys.argv[1] == 'manifest':
        source_manifest(sys.argv[2])
    elif sys.argv[1] == 'storage':
        storage_manifest(sys.argv[2])
    elif sys.argv[1] == 'get':
        private.renew()
        result = api.call(sys.argv[2])
        save(sys.argv[3], result)
        print(json.dumps({'saved': sys.argv[3], 'type': type(result).__name__}))
    elif sys.argv[1] == 'scaffold':
        scaffold()
    elif sys.argv[1] == 'verify':
        verify()
    elif sys.argv[1] == 'verify-rev10':
        verify_rev10()
    elif sys.argv[1] == 'documents':
        make_test_documents()
    elif sys.argv[1] == 'intake':
        intake()
    elif sys.argv[1] == 'boq':
        boq()
    elif sys.argv[1] == 'outbound-prepare':
        outbound_prepare()
    elif sys.argv[1] == 'outbound-finish':
        outbound_finish()
    elif sys.argv[1] == 'input-chain':
        input_chain()
    elif sys.argv[1] == 'inspection-assets':
        make_inspection_assets()
    elif sys.argv[1] == 'planning':
        planning()
    elif sys.argv[1] == 'activate-planning':
        activate_planning()
    elif sys.argv[1] == 'execution-photo':
        prepare_execution_photo()
    elif sys.argv[1] == 'correct-test-schedule':
        correct_test_schedule()
    elif sys.argv[1] == 'finish-execution':
        finish_execution()
    elif sys.argv[1] == 'client-prepare':
        prepare_client()
    elif sys.argv[1] == 'client-snapshot':
        client_snapshot(sys.argv[2])
