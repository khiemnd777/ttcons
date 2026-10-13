"""Read-only audit of actual BM06/BM07 downloads, Google readbacks and baselines."""
import hashlib
import io
import json
import posixpath
from pathlib import Path
import xml.etree.ElementTree as ET
import zipfile
import openpyxl
from PIL import Image

OUT = Path(__file__).resolve().parent
ROOT = OUT.parents[2]
CLIENT = Path('/Users/khiemnguyen/Works/ttcons/documents/Bo_Bieu_Mau_Giam_Sat_Thu_Thiem.xlsx')
S = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
D = '{http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing}'
A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
R = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'


def load(name):
    return json.loads((OUT / name).read_text())


def verify():
    checks = {}
    def check(name, ok):
        checks[name] = bool(ok)
    books = {n: openpyxl.load_workbook(OUT / (n + '.xlsx')) for n in ['BM06-form', 'BM06-daily', 'BM07']}
    client = openpyxl.load_workbook(CLIENT)
    info, parity, errors = {}, {}, []
    for n, book in books.items():
        prefix = n.lower()
        raw = (OUT / (n + '.xlsx')).read_bytes()
        check(prefix + '_actual_download_bytes', raw == Path(load(prefix + '-download.json')['sourcePath']).read_bytes())
        template = 'BM-GS-07' if n == 'BM07' else 'BM-GS-06'
        sheet = book.worksheets[0]
        count = 7 if n == 'BM07' else 10
        check(prefix + '_original_headers', all(sheet.cell(7, c).value == client[template].cell(7, c).value for c in range(1, count + 1)))
        check(prefix + '_project_header', ('B5:G5' if n == 'BM07' else 'B5:J5') in {str(m) for m in sheet.merged_cells} and sheet.row_dimensions[5].height == 33)
        imp, drive, meta = [load(prefix + '-google-' + suffix + '.json')['structuredContent'] for suffix in ['import', 'drive', 'metadata']]
        check(prefix + '_private_native_google', imp['converted'] and imp['mimeType'] == 'application/vnd.google-apps.spreadsheet' and drive['shared'] is False and all(p['role'] == 'owner' and p['type'] == 'user' for p in drive['permissions']))
        check(prefix + '_same_tab_names', [s['properties']['title'] for s in meta['sheets']] == book.sheetnames)
        cells = {}
        for tab in load(prefix + '-google-cells.json')['structuredContent']['sheets']:
            name = tab['properties']['title']
            for block in tab.get('data', []):
                for r, row in enumerate(block.get('rowData', []), block.get('startRow', 0) + 1):
                    for c, value in enumerate(row.get('values', []), block.get('startColumn', 0) + 1):
                        entered = value.get('userEnteredValue', {})
                        v = entered.get('formulaValue', entered.get('stringValue', entered.get('numberValue', entered.get('boolValue'))))
                        cells[(name, r, c)] = v
                        if value.get('effectiveValue', {}).get('errorValue') or value.get('formattedValue', '').startswith(('#VALUE!', '#REF!', '#ERROR!', '#NUM!', '#N/A', '#NAME?', '#DIV/0!')):
                            errors.append((n, name, r, c, value))
        mismatch = []
        compared = 0
        for s in book:
            rows, cols = (16, 7) if n == 'BM07' else ((15, 10) if s == sheet else (25, 8))
            for r in range(1, rows + 1):
                for c in range(1, cols + 1):
                    # Native image values are verified through XML + rendered cell menus.
                    if n == 'BM07' and c == 6 and 8 <= r <= 10:
                        continue
                    compared += 1
                    expected = s.cell(r, c).value
                    if cells.get((s.title, r, c)) != expected:
                        mismatch.append([s.title, r, c, expected, cells.get((s.title, r, c))])
        parity[n] = {'compared': compared, 'mismatches': mismatch}
        check(prefix + '_google_value_parity', not mismatch)
        check(prefix + '_google_header_selection_100', ('Value: B5:G5' if n == 'BM07' else 'Value: B5:J5') in (OUT / (prefix + '-google-header-state.txt')).read_text() and '100% selected.' in (OUT / (prefix + '-google-header-state.txt')).read_text())
        info[n] = {'bytes': len(raw), 'sha256': hashlib.sha256(raw).hexdigest(), 'tabs': book.sheetnames, 'row5_height': sheet.row_dimensions[5].height}
    check('google_no_effective_errors', not errors)
    form = load('bm06-source-form-before.json')
    row = form['answers_snapshot_json']['next_plan_items'][0]
    events = load('bm06-form-events-before.json')
    facts = next(e['metadata_json']['native_daily_sources'] for e in events if e['event_type'] == 'amended' and e['metadata_json']['native_daily_sources']['form_revision_no'] == form['current_revision'])
    frozen = facts['structured_json']['next_plan_items'][0]
    expected = [1, row['instruction_code'], row['title'], row['location'], None, row['receiver'] + ' • ' + row['contact'], None, row['due_date'], 0, frozen['state']]
    check('bm06_exact_submitted_v5_r2', form['status'] == 'submitted' and form['current_revision'] == 2 and load('bm06-source-version.json')['version_no'] == 5)
    check('bm06_frozen_state_exact_task_identity', frozen['state_source']['source_id'] == load('bm06-generated-tasks-before.json')[0]['task']['id'] and frozen['state_source']['status'] == 'in_progress' and frozen['row_id'] == row['row_id'])
    check('bm06_source_missing_facts_not_invented', 'subcontractor' not in row and 'issued_date' not in row)
    for n in ['BM06-form', 'BM06-daily']:
        check(n.lower() + '_exact_ten_values_real_zero', [books[n].active.cell(8, c).value for c in range(1, 11)] == expected)
        check(n.lower() + '_signatures', all(books[n].active[a].value == client['BM-GS-06'][a].value for a in ['A12', 'D12', 'G12', 'A13', 'D13', 'G13']))
    photos = load('bm07-source-photos-before.json')
    sheet = books['BM07'].active
    check('bm07_exact_three_photo_rows_one_tab', len(photos) == 3 and len(books['BM07'].sheetnames) == 1 and [sheet.cell(r, 1).value for r in range(8, 11)] == [1, 2, 3])
    check('bm07_unknown_capture_time_remains_blank', all(not p['payload_json']['capture_metadata'].get('captured_at') for p in photos) and all(sheet.cell(r, 3).value is None for r in range(8, 11)))
    check('bm07_source_row_metadata', all(sheet.cell(i + 8, 2).value == p['payload_json']['capture_metadata']['location'] and sheet.cell(i + 8, 4).value == (p['caption'] or p['description']) and sheet.cell(i + 8, 5).value == {'defect': 'Phát hiện lỗi', 'during': 'Trong thi công'}[p['payload_json']['capture_metadata']['photo_phase']] for i, p in enumerate(photos)))
    check('bm07_signatures', all(sheet[a].value == client['BM-GS-07'][a].value for a in ['A13', 'D13', 'A14', 'D14']))
    package = zipfile.ZipFile(OUT / 'BM07.xlsx')
    drawing = ET.fromstring(package.read('xl/drawings/drawing1.xml'))
    rels = ET.fromstring(package.read('xl/drawings/_rels/drawing1.xml.rels'))
    anchors = drawing.findall(D + 'oneCellAnchor')
    media_info = []
    for i, anchor in enumerate(anchors):
        p = photos[i]
        origin = anchor.find(D + 'from')
        rid = anchor.find('.//' + A + 'blip').attrib[R + 'embed']
        target = next(r.attrib['Target'] for r in rels if r.attrib['Id'] == rid)
        part = target.lstrip('/') if target.startswith('/') else posixpath.normpath(posixpath.join('xl/drawings', target))
        raw = package.read(part)
        image = Image.open(io.BytesIO(raw)); image.load()
        original = Image.open(ROOT / 'storage' / p['storage_key']); original.load()
        ext = anchor.find(D + 'ext'); cx, cy = int(ext.attrib['cx']), int(ext.attrib['cy'])
        descr = anchor.find('.//' + D + 'cNvPr').attrib['descr']
        check('bm07_F' + str(i + 8) + '_identity_and_fit', origin.find(D + 'col').text == '5' and int(origin.find(D + 'row').text) == i + 7 and origin.find(D + 'colOff').text == origin.find(D + 'rowOff').text == '0' and p['id'] in descr and cx <= 168 * 9525 and cy <= sheet.row_dimensions[i + 8].height * 12700 and abs(cx / cy - image.width / image.height) < 0.00001)
        check('bm07_F' + str(i + 8) + '_native_cell_ui', 'Value: F' + str(i + 8) in (OUT / ('bm07-google-f' + str(i + 8) + '-cell-state.txt')).read_text() and 'Put image over cells' in (OUT / ('bm07-google-f' + str(i + 8) + '-cell-state.txt')).read_text())
        check('bm07_F' + str(i + 8) + '_original_checksum_and_derivative', hashlib.sha256((ROOT / 'storage' / p['storage_key']).read_bytes()).hexdigest() == p['checksum_sha256'] and max(image.size) <= 1280 and abs(image.width / image.height - original.width / original.height) < 0.002)
        if image.size == original.size:
            check('bm07_F' + str(i + 8) + '_original_pixels', image.convert('RGB').tobytes() == original.convert('RGB').tobytes())
        media_info.append({'cell': 'F' + str(i + 8), 'evidence_id': p['id'], 'media_part': part, 'media_sha256': hashlib.sha256(raw).hexdigest(), 'bytes': len(raw), 'size': image.size, 'original_size': original.size, 'extent_emu': [cx, cy]})
    check('bm07_three_only_drawings_and_budget', len(anchors) == 3 and len(drawing) == 3 and sum(m['bytes'] for m in media_info) <= 100 * 1024 ** 2 and (OUT / 'BM07.xlsx').stat().st_size <= 150 * 1024 ** 2)
    before = {**load('before-database.json'), **load('before-construction-history.json')}
    after = {**load('after-database.json'), **load('after-construction-history.json')}
    changed, additions, checked = [], [], 0
    for table, rows in before.items():
        if not isinstance(rows, list):
            continue
        old, now = [{r['id']: r['md5'] for r in x} for x in [rows, after[table]]]
        checked += len(old)
        changed.extend((table, key) for key in old if old[key] != now.get(key))
        additions.extend((table, key) for key in now if key not in old)
    audit = load('new-audit-rows.json')
    check('all_old_rows_retained_unchanged', not changed)
    check('only_four_actual_export_audits_added', len(additions) == len(audit) == 4 and all(t == 'audit_logs' for t, _ in additions) and all(r['event_type'] == 'native_source_exported' for r in audit))
    check('bm06_form_and_events_unchanged', form == load('bm06-source-form-after.json') and events == load('bm06-form-events-after.json'))
    check('bm07_all_source_evidence_unchanged', photos == load('bm07-source-photos-after.json'))
    check('all_storage_unchanged', load('storage-before.json') == load('storage-after.json') and len(load('storage-before.json')) == 149)
    check('same_database_migrations158', before['database'] == after['database'] == 'consflow' and before['migration_count'] == after['migration_count'] == 158)
    built = load('source-build-manifest.json')['files']
    actual = {n: hashlib.sha256((ROOT / n).read_bytes()).hexdigest() for n in built}
    delta = [n for n, v in load('source-before.json')['files'].items() if actual[n] != v]
    check('all979_current_hashes_match_build', actual == built and len(actual) == 979)
    check('only_approved_renderer_and_test_changed', set(delta) == {'backend/internal/core/bm_gs_workbook.go', 'backend/internal/core/bm_gs_risk_header_test.go'})
    runtime = load('runtime.json')
    check('only_backend_recreated_same_development', runtime['only_backend_recreated'] and runtime['backend_image_changed'] and runtime['backend']['scope_env']['APP_ENV'] == 'development')
    check('build_update_readiness', runtime['build_exit_code'] == runtime['update_exit_code'] == 0 and runtime['readiness_http'] == runtime['frontend_http'] == 200)
    return {'passed': all(checks.values()), 'checks': checks, 'files': info, 'google_parity': parity, 'google_errors': errors, 'images': media_info, 'old_unique_rows_checked': checked, 'changed_or_lost_rows': changed, 'new_rows': additions, 'source_delta': delta, 'limits': 'Development Admin; BM06 pinned v5 submitted R2 one next-plan row, BM07 three landscape synthetic imported files. User UX signoff pending. No actual camera/device/portrait/square/tall/capacity/other roles/Excel desktop/print rerun. Audit source revision_no=0 denotes the current-source export contract; form R2 is separately proven from immutable events. Four audits include one BM07 repeat for download retention.'}


if __name__ == '__main__':
    result = verify()
    (OUT / 'verification.json').write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n')
    print(json.dumps(result, ensure_ascii=False, indent=2))
    assert result['passed']
