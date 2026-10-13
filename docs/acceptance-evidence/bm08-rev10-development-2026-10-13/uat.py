"""BM08 authorized full Rev10 UAT: canonical API writes, read-only SQL evidence.

Never changes historical fixtures, applied migrations, exports or Google cells.
"""
import datetime
import hashlib
import importlib.util
import json
from pathlib import Path
import subprocess
import sys

ROOT = Path(__file__).resolve().parents[3]
OUT = Path(__file__).resolve().parent
PRIOR = ROOT / 'docs/acceptance-evidence/bm03-client-reconciliation-development-2026-10-13'
spec = importlib.util.spec_from_file_location('bm03_uat', PRIOR / 'uat.py')
parent = importlib.util.module_from_spec(spec)
spec.loader.exec_module(parent)
api = parent.api
parent.OUT = OUT


def save(name, value):
    (OUT / name).write_text(json.dumps(value, ensure_ascii=False, indent=2) + '\n')


def state():
    path = OUT / 'fixture-state.json'
    if not path.exists():
        save('fixture-state.json', json.loads((PRIOR / 'fixture-state.json').read_text()))
    return json.loads(path.read_text())


def baseline(prefix):
    tables = ['projects', 'project_members', 'work_items', 'pipeline_templates', 'pipeline_template_nodes', 'pipeline_template_steps', 'pipeline_template_edges', 'pipeline_runs', 'pipeline_run_nodes', 'pipeline_run_steps', 'workflow_form_templates', 'workflow_form_versions', 'workflow_form_fields', 'workflow_form_instances', 'workflow_form_answers', 'workflow_form_events', 'workflow_form_signatures', 'workflow_form_pdf_exports', 'acceptance_records', 'acceptance_checklist_items', 'handover_records', 'handover_checklist_items', 'documents', 'subcontractor_reconciliations', 'subcontractor_reconciliation_items', 'owner_client_reconciliations', 'owner_client_reconciliation_items', 'evidence_records', 'work_reports', 'work_report_revisions', 'work_report_events', 'incoming_acceptance_records', 'incoming_acceptance_items', 'stock_movements']
    existing = parent.sql("SELECT coalesce(json_agg(table_name),'[]') FROM information_schema.columns WHERE table_schema='public' AND column_name='id'")
    tables = sorted(set(t for t in tables if t in existing) | {t for t in existing if t.startswith(('pipeline_', 'workflow_form_', 'subcontractor_checklist_', 'owner_client_checklist_', 'construction_handoff'))})
    parts = ["'" + t + "',(SELECT coalesce(json_agg(json_build_object('id',id,'md5',md5(row_to_json(r)::text)) ORDER BY id),'[]') FROM " + t + ' r)' for t in tables]
    data = parent.sql("SELECT json_build_object('database',current_database(),'migration_count',(SELECT count(*) FROM schema_migrations)," + ','.join(parts) + ')')
    save(prefix + '-database.json', data)
    print(json.dumps({'prefix': prefix, 'tables': len(tables), 'rows': sum(len(data[t]) for t in tables), 'migrations': data['migration_count']}))


def inventory(prefix):
    parent.private.renew()
    st = state()
    pid = st['project_id']
    for suffix in ['pipeline-run', 'work-items', 'handover/options', 'acceptances', 'incoming-acceptances']:
        save(prefix + '-' + suffix.replace('/', '-') + '.json', api.call('/projects/' + pid + '/' + suffix))
    items = api.call('/projects/' + pid + '/work-items')
    print(json.dumps([{k: w.get(k) for k in ['id', 'title', 'status', 'run_node_business_code', 'run_step_business_code', 'workflow_form_template_code', 'domain_action_key']} for w in items if w['run_node_business_code'] in ['incoming_acceptance', 'outgoing_acceptance', 'handover_finance']], ensure_ascii=False))
    if prefix == 'before':
        mutable = {'projects': [pid], 'pipeline_runs': [st['run_id']], 'pipeline_run_nodes': [], 'pipeline_run_steps': [], 'work_items': [], 'workflow_form_instances': []}
        run = api.call('/projects/' + pid + '/pipeline-run')
        for node in run['nodes']:
            if node['status'] not in ['completed', 'skipped']:
                mutable['pipeline_run_nodes'].append(node['id'])
                mutable['pipeline_run_steps'].extend(step['id'] for step in node['steps'])
        for item in items:
            if item['status'] not in ['completed', 'cancelled']:
                mutable['work_items'].append(item['id'])
                if item.get('workflow_form_instance_id'):
                    mutable['workflow_form_instances'].append(item['workflow_form_instance_id'])
        save('authorized-mutable-source-ids.json', mutable)


def incoming_and_final_draft():
    parent.private.renew()
    st = state()
    incoming = 'b6cdcf60-a200-4ecf-a306-53f25a110733'
    work = api.call('/work-items/' + incoming)
    if work['status'] == 'ready':
        api.call('/work-items/' + incoming + '/start', 'POST', {})
    name = '[Synthetic] Waterproof test floor'
    parent.submit_form(incoming, {'material_items': [{'material': name, 'delivered_quantity': 25, 'accepted_quantity': 25, 'unit': 'm2', 'decision': 'accepted'}]}, 'incoming-form')
    if not st.get('incoming_acceptance_id'):
        # The approved BOQ line is work scope, without a material identity/name.
        # Use the supported supplier receipt source rather than inventing a PO-line identity.
        request = {'work_item_id': incoming, 'source_type': 'supplier', 'supplier_partner_id': st['supplier_partner_id'], 'received_at': '2026-10-13', 'delivery_reference': 'UAT-BM08-REV10-20261013', 'comment': '[Synthetic] Explicit supplier receipt test source. No procurement-line/physical delivery/implicit stock claim.', 'items': [{'item_index': 0, 'material_name': name, 'unit': 'm2', 'delivered_quantity': 25, 'accepted_quantity': 25, 'rejected_quantity': 0, 'decision': 'accepted', 'condition': 'usable', 'comment': '[Synthetic] Explicit named test input; no invented BOQ material relation.'}]}
        save('incoming-request.json', request)
        record = api.call('/projects/' + st['project_id'] + '/incoming-acceptances', 'POST', request)
        st['incoming_acceptance_id'] = record['id']
        save('fixture-state.json', st)
        save('incoming-created.json', record)
    record = api.call('/incoming-acceptances/' + st['incoming_acceptance_id'])
    if record['status'] != 'submitted':
        save('incoming-submitted.json', api.call('/incoming-acceptances/' + st['incoming_acceptance_id'] + '/submit', 'POST', {}))
    parent.complete_work(incoming, 'incoming-material')
    parent.advance_current('incoming_acceptance')
    final = '4fad1dc4-13b5-46b8-982b-f07ee3d875c1'
    work = api.call('/work-items/' + final)
    if work['status'] == 'ready':
        api.call('/work-items/' + final + '/start', 'POST', {})
    form = parent.form_for(final)
    st['final_form_id'] = form['instance']['id']
    save('fixture-state.json', st)
    answers = {'check_date': '2026-10-13', 'client_representative': '[Synthetic] Đại diện CĐT UAT', 'company_representative': 'Consflow Admin', 'handover_scope': 'handover_item', 'scope_detail': '[Synthetic] Hồ sơ bàn giao vùng thử Rev10', 'checklist_items': [{'title': '[Synthetic] Hồ sơ vùng thử theo nguồn đã chốt', 'output_ref': 'UAT-BM08-REV10', 'status': 'done', 'client_decision': 'pass', 'comment': '[Synthetic] Chỉ kiểm luồng hồ sơ nội bộ, không xác nhận chất lượng công trình thật.'}], 'final_client_note': '[Synthetic] Không có hiệu lực pháp lý; dữ liệu nghiệm thu Development.'}
    if form['instance']['status'] == 'draft':
        save('final-draft.json', api.call('/workflow-form-instances/' + st['final_form_id'] + '/draft', 'PUT', {'answers': answers, 'expected_answers': form['answers'], 'comment': '[Synthetic] Prepare exact pinned form for actual UI chốt/signature/PDF.'}))
    save('final-source-before-ui.json', parent.form_for(final))
    print(json.dumps({'incoming_acceptance': st['incoming_acceptance_id'], 'final_work_item': final, 'final_form': st['final_form_id']}))


def final_acceptance_and_advance():
    parent.private.renew()
    st = state()
    wid = '4fad1dc4-13b5-46b8-982b-f07ee3d875c1'
    form = parent.form_for(wid)
    save('final-signed-pdf-source.json', form)
    assert form['instance']['status'] == 'submitted' and len(form['signatures']) == 1 and len(form['pdf_exports']) == 1
    pdf = form['pdf_exports'][0]
    st.update(final_pdf_id=pdf['id'], final_signature_id=form['signatures'][0]['id'])
    save('fixture-state.json', st)
    if not st.get('final_acceptance_id'):
        request = {'work_item_id': wid, 'form_instance_id': st['final_form_id'], 'pdf_export_id': st['final_pdf_id'], 'type': 'final', 'title': '[Synthetic] Nghiệm thu cuối BM08 Rev10', 'acceptance_date': '2026-10-13', 'location': '[Synthetic] Vùng thử nội bộ Rev10', 'client_representative': '[Synthetic] Đại diện CĐT UAT', 'contractor_representative': 'Consflow Admin', 'scope_summary': '[Synthetic] Bàn giao hồ sơ kiểm thử theo quy trình đầy đủ', 'notes': '[Synthetic] Nghiệm thu kỹ thuật luồng Development, không có hiệu lực pháp lý.', 'checklist': [{'item_index': 0, 'title': '[Synthetic] Hồ sơ vùng thử theo nguồn đã chốt', 'description': '[Synthetic] Có checklist đã chốt, xác nhận UAT và PDF bất biến của đúng nguồn.', 'decision': 'pass', 'severity': 'low', 'photo_required': False, 'note': '[Synthetic] Đã kiểm tra hồ sơ thử nội bộ'}]}
        save('final-acceptance-request.json', request)
        record = api.call('/projects/' + st['project_id'] + '/acceptances', 'POST', request)
        st['final_acceptance_id'] = record['id']; save('fixture-state.json', st)
        save('final-acceptance-created.json', record)
    record = api.call('/acceptances/' + st['final_acceptance_id'])
    assert record['signature_status'] == 'signed' and record['document_status'] == 'linked'
    if record['status'] == 'draft':
        save('final-acceptance-submitted.json', api.call('/acceptances/' + record['id'] + '/submit', 'POST', {}))
    record = api.call('/acceptances/' + st['final_acceptance_id'])
    if record['status'] != 'completed':
        save('final-acceptance-passed.json', api.call('/acceptances/' + record['id'] + '/pass', 'POST', {'comment': '[Synthetic] Hồ sơ thử có đúng nguồn ký/PDF và tiêu chí đã kiểm; không xác nhận chất lượng công trình thật.'}))
    reconciliation = api.call('/work-items/' + wid + '/owner-client-reconciliation')
    save('owner-reconciliation-before.json', reconciliation)
    if reconciliation['reconciliation']['status'] != 'submitted':
        request = {'comment': '[Synthetic] Đối chiếu đúng nghiệm thu cuối cùng công trình đã pass', 'items': [{'item_index': 0, 'acceptance_id': st['final_acceptance_id'], 'decision': 'accepted', 'comment': '[Synthetic] Nguồn hồ sơ thử hợp lệ'}]}
        save('owner-reconciliation-submitted.json', api.call('/work-items/' + wid + '/owner-client-reconciliation/submit', 'POST', request))
    parent.complete_work(wid, 'final-package')
    parent.advance_current('outgoing_acceptance')
    save('handover-options-valid.json', api.call('/projects/' + st['project_id'] + '/handover/options'))
    print(json.dumps({'final_acceptance_id': st['final_acceptance_id'], 'pdf_id': st['final_pdf_id']}))


DOCUMENTS = {'as_built_documents': 'As-built test dossier', 'operation_manuals': 'Operation and maintenance test instructions', 'warranty_documents': 'Warranty test dossier - no commercial commitment', 'signed_handover_minutes': 'Two-party internal test handover acknowledgement'}


def documents():
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet
    for key, title in DOCUMENTS.items():
        file = OUT / (key + '.pdf')
        if file.exists():
            continue
        styles = getSampleStyleSheet()
        story = [Paragraph('SYNTHETIC DEVELOPMENT UAT', styles['Title']), Spacer(1, 18), Paragraph(title, styles['Heading1']), Spacer(1, 18), Paragraph('Project BM03-REV10-20261013 / 432115cb-bd5b-4fae-9edc-74ce031eacf3. This is an internal test source for the existing full Rev10 handover. Not a real construction, client, warranty, or legal document.', styles['BodyText']), Spacer(1, 18), Paragraph('Source: completed test final acceptance with the exact submitted owner form, internal UAT signature and native PDF. Planned test handover date: 14 October 2026. Actual physical quality, delivery and client consent are outside this test.', styles['BodyText']), Spacer(1, 18)]
        content = {'as_built_documents': 'Test scope: named25m2 waterproof floor. Source drawing reference UAT-ASBUILT-R10. The record identifies the test area and existing acceptance source; it is not a physical survey.', 'operation_manuals': 'Internal test review: confirm dossier source identity and file readability. No actual site operation or safety instruction is issued.', 'warranty_documents': 'Internal test warranty dossier identifier UAT-WARRANTY-R10. No payment, guarantee or warranty period is legally agreed.', 'signed_handover_minutes': 'Both parties are simulated: Synthetic client representative and Consflow Admin. Internal acknowledgement marks UAT CLIENT / UAT CONTRACTOR, no legal effect. The real application test signature/PDF belongs to the same final source.'}[key]
        story.extend([Paragraph(content, styles['BodyText']), Spacer(1, 26), Paragraph('UAT REVIEWED SOURCE / QA TEST ONLY', styles['Heading2'])])
        SimpleDocTemplate(str(file), pagesize=A4, leftMargin=44, rightMargin=44, topMargin=44, bottomMargin=44).build(story)
    st = state()
    parent.private.renew()
    parent.download_protected('/workflow-form-pdf-exports/' + st['final_pdf_id'] + '/file', OUT / 'final-native-signed.pdf')
    print('Created four synthetic documents and downloaded native immutable final PDF; inspect before uploads.')


def upload_documents():
    parent.private.renew()
    st = state()
    st.setdefault('handover_documents', {})
    for key in DOCUMENTS:
        if key in st['handover_documents']:
            continue
        record = parent.upload_file('/projects/' + st['project_id'] + '/documents/upload', OUT / (key + '.pdf'), {'document_type': key, 'name': '[Synthetic] BM08 Rev10 ' + key})
        st['handover_documents'][key] = record['id']
        save('fixture-state.json', st)
        save(key + '-uploaded.json', record)
    print('Uploaded four same-project actual PDF files; checklist links must be selected via UI.')


def handover_snapshot(label):
    parent.private.renew()
    st = state()
    rows = api.call('/projects/' + st['project_id'] + '/handover')
    save(label + '-list.json', rows)
    current = next(h for h in rows if h['handover_code'] == 'UAT-BM08-REV10-20261013')
    st['handover_id'] = current['id']; save('fixture-state.json', st)
    record = api.call('/handover/' + st['handover_id'])
    save(label + '-detail.json', record)
    dates = parent.sql("SELECT json_build_object('business_date',handover_date::text,'server_current_date',CURRENT_DATE::text,'status',status,'final_acceptance_id',final_acceptance_id,'completed_at',completed_at) FROM handover_records WHERE id='" + st['handover_id'] + "'")
    save(label + '-sql-date.json', dates)
    assert dates['business_date'] == '2026-10-14'
    save(label + '-physical-work-item.json', api.call('/work-items/73e77a9e-7abd-4f2c-b768-c62e3fd31627'))
    save(label + '-finance-work-item.json', api.call('/work-items/d787a8c0-cc4c-442e-bf36-6242f25d65ea'))
    print(json.dumps({'id': st['handover_id'], 'date': dates['business_date'], 'server_date': dates['server_current_date'], 'status': dates['status'], 'detail_keys': list(record)}, ensure_ascii=False))


def verify():
    import zipfile
    import openpyxl
    load = lambda name: json.loads((OUT / name).read_text())
    before = load('before-database.json')
    after = load('after-database.json')
    allowed = load('authorized-mutable-source-ids.json')
    changes, expected, checked = [], [], 0
    for table, rows in before.items():
        if not isinstance(rows, list):
            continue
        current = {row['id']: row['md5'] for row in after[table]}
        for row in rows:
            checked += 1
            if current.get(row['id']) != row['md5']:
                target = expected if row['id'] in allowed.get(table, []) else changes
                target.append({'table': table, 'id': row['id']})
    old_files, new_files = load('storage-before.json'), load('storage-after.json')
    file_changes = [name for name, checksum in old_files.items() if new_files.get(name) != checksum]
    source = load('source-before.json')['files']
    source_changes = [name for name, checksum in source.items() if hashlib.sha256((ROOT / name).read_bytes()).hexdigest() != checksum]
    saved, ready, final = [load(label + '-detail.json') for label in ['after-create', 'before-complete', 'final']]
    st = state()
    dates = [load(label + '-sql-date.json') for label in ['after-create', 'before-complete', 'final']]
    assertions = {
        'date_saved_before_complete': saved['handover_date'][:10] == '2026-10-14',
        'date_preserved_through_complete': all(d['business_date'] == '2026-10-14' for d in dates),
        'entered_date_differs_from_server_default': dates[0]['business_date'] != dates[0]['server_current_date'],
        'valid_own_final_acceptance': final['final_acceptance_id'] == st['final_acceptance_id'] and final['final_acceptance_status'] == 'completed' and final['final_acceptance_result'] == 'pass',
        'handover_ready_before_complete': ready['status'] == 'ready' and ready['documents_ready'] and ready['warehouse_closed'],
        'handover_completed': final['status'] == 'handed_over' and final['completed_at'] is not None,
        'six_required_verified': len(final['checklist_items']) == 6 and all(i['required'] and i['status'] == 'verified' for i in final['checklist_items']),
        'six_explicit_assessments': all(i['physical_result'] == 'pass' for i in final['checklist_items']),
        'conclusion_preserved': final['handover_conclusion'] == 'eligible',
        'history_unchanged_outside_authorized_runtime_sources': not changes,
        'old_storage_files_unchanged': not file_changes,
        'source_unchanged_no_build_needed': not source_changes,
        'no_migration': before['migration_count'] == after['migration_count'] == 158,
        'physical_work_item_left_in_progress': load('final-physical-work-item.json')['status'] == 'in_progress',
        'accounting_not_started_or_completed': load('final-finance-work-item.json')['status'] in ['pending', 'ready'] and load('final-finance-work-item.json')['started_at'] is None and load('final-finance-work-item.json')['completed_at'] is None,
    }
    for key, doc_id in st['handover_documents'].items():
        item = next(i for i in final['checklist_items'] if i['item_code'] == key)
        assertions[key + '_exact_source'] = item['document_id'] == doc_id
        document = load(key + '-uploaded.json')
        assertions[key + '_same_project'] = document['project_id'] == st['project_id']
    file = OUT / 'BM08-Rev10-handed-over.xlsx'
    native = Path('/Users/khiemnguyen/Downloads/BM-GS-08 (8).xlsx')
    assertions['actual_download_byte_equal'] = file.read_bytes() == native.read_bytes()
    archive = zipfile.ZipFile(file)
    assertions['zip_valid'] = archive.testzip() is None
    workbook = openpyxl.load_workbook(file)
    sheet = workbook.worksheets[0]
    assertions['worksheet_business_date'] = sheet.title == '2026-10-14_01'
    assertions['physical_assessments_not_verification_codes'] = all(sheet.cell(i, 4).value == 'Đạt' for i in range(8, 14))
    assertions['no_fabricated_defect_deadline'] = all(sheet.cell(i, 5).value in [None, ''] and sheet.cell(i, 6).value in [None, ''] for i in range(8, 14))
    assertions['client_header_columns_preserved'] = [sheet.cell(7, c).value for c in range(1, 7)] == ['STT', 'DANH MỤC HẠNG MỤC BÀN GIAO', 'YÊU CẦU HỒ SƠ & THỰC TẾ', 'ĐÁNH GIÁ (ĐẠT/KĐ)', 'TỒN TẠI CẦN XỬ LÝ (NẾU CÓ)', 'THỜI HẠN XỬ LÝ']
    google = load('google-cells.json')['structuredContent']
    assertions['google_worksheet_date_preserved'] = google['sheets'][0]['properties']['title'] == sheet.title
    data = google['sheets'][0]['data'][0].get('rowData', [])
    mismatch, errors, populated = [], [], 0
    for i in range(32):
        cells = data[i].get('values', []) if i < len(data) else []
        for j in range(6):
            cell = cells[j] if j < len(cells) else {}
            effective = cell.get('effectiveValue', {})
            actual = effective.get('stringValue', effective.get('numberValue', effective.get('boolValue')))
            value = sheet.cell(i+1, j+1).value
            value = None if value in [None, ''] else value
            populated += value is not None
            if actual != value:
                mismatch.append([i+1, j+1, value, actual])
            if 'errorValue' in effective or str(cell.get('formattedValue','')).startswith('#VALUE'):
                errors.append([i+1, j+1, cell])
    assertions['google_all_values_match'] = not mismatch
    assertions['google_no_effective_error'] = not errors
    result = {'assertions': assertions, 'passed': all(assertions.values()), 'old_rows_checked': checked, 'unexpected_old_row_changes': changes, 'expected_canonical_source_changes': expected, 'old_files_checked': len(old_files), 'old_file_changes': file_changes, 'source_files_checked': len(source), 'source_changes': source_changes, 'google_cells_checked': 192, 'google_populated_cells': populated, 'google_mismatches': mismatch, 'google_errors': errors, 'download_sha256': hashlib.sha256(file.read_bytes()).hexdigest(), 'scope': 'Synthetic Admin full published Rev10 source path to handed_over BM08/date/export/Google; separate physical WI completion rejected by automatic approval review and remains in_progress, finance not completed, user UX acceptance pending.'}
    save('evidence-verification.json', result)
    print(json.dumps(result, ensure_ascii=False))
    assert result['passed']


if __name__ == '__main__':
    command = sys.argv[1]
    if command == 'baseline':
        baseline(sys.argv[2])
    elif command == 'inventory':
        inventory(sys.argv[2])
    elif command == 'manifest':
        parent.source_manifest(sys.argv[2])
    elif command == 'storage':
        parent.storage_manifest(sys.argv[2])
    elif command == 'get':
        parent.private.renew()
        save(sys.argv[3], api.call(sys.argv[2]))
        print('Saved', sys.argv[3])
    elif command == 'incoming-final':
        incoming_and_final_draft()
    elif command == 'final-advance':
        final_acceptance_and_advance()
    elif command == 'documents':
        documents()
    elif command == 'upload-documents':
        upload_documents()
    elif command == 'handover-snapshot':
        handover_snapshot(sys.argv[2])
    elif command == 'verify':
        verify()
