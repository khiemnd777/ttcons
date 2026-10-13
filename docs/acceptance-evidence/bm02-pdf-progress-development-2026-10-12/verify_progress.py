"""Evidence checks for the approved Development PDF progress presentation fix.

Reuses the previous read-only business baseline helper and its temporary auth
session. This writes local evidence only, not business data, PDFs or XLSX.
"""
import datetime
import hashlib
import importlib.util
import json
from pathlib import Path
import re
import subprocess
import sys
from pypdf import PdfReader
from openpyxl import load_workbook
from zipfile import ZipFile

ROOT = Path(__file__).resolve().parents[3]
OUT = Path(__file__).resolve().parent
WEATHER = ROOT / 'docs/acceptance-evidence/bm02-weather-development-2026-10-12'
spec = importlib.util.spec_from_file_location('weather_evidence', WEATHER / 'verify_weather.py')
previous = importlib.util.module_from_spec(spec)
spec.loader.exec_module(previous)
previous.OUT = OUT


def save(name, value):
    (OUT / name).write_text(json.dumps(value, ensure_ascii=False, indent=2) + '\n')


def baseline(prefix):
    previous.records(prefix)
    manual_file = ROOT / 'docs/acceptance-evidence/bm02-delivery-revision-development-2026-10-12/fixture-state.json'
    if manual_file.exists():
        value = json.loads(manual_file.read_text())
        record = value.get('data', value)
        if isinstance(record.get('report'), dict):
            record = record['report']
        source_id = record.get('source_id')
        if source_id:
            save(prefix + '-manual-report.json', previous.uat.api.call('/work-reports/daily-reports/' + source_id))


def manifest():
    base = json.loads((WEATHER / 'source-manifest.json').read_text())
    extras = json.loads((OUT / 'extra-source-files.json').read_text()) if (OUT / 'extra-source-files.json').exists() else []
    names = set(base) | set(extras)
    current = {name: hashlib.sha256((ROOT / name).read_bytes()).hexdigest() for name in sorted(names)}
    save('source-baseline-manifest.json', base)
    save('source-manifest.json', current)
    save('source-state.json', {
        'captured_at': datetime.datetime.now(datetime.timezone.utc).isoformat(),
        'head': subprocess.check_output(['git', 'rev-parse', 'HEAD'], cwd=ROOT, text=True).strip(),
        'file_count': len(current),
        'manifest_sha256': hashlib.sha256((OUT / 'source-manifest.json').read_bytes()).hexdigest(),
        'changed_existing_files': [name for name in base if base[name] != current[name]],
        'new_files': sorted(set(current) - set(base)),
    })


def verify():
    baseline('after')
    read = lambda name: json.loads((OUT / name).read_text())
    result = {'scope': 'Development selected-revision PDF recorded progress; source/snapshot data preserved, existing export audit retained', 'checks': {}, 'files': {}}

    def check(name, value):
        result['checks'][name] = bool(value)
        assert value, name

    for name in ['database', 'form', 'report', 'receipt', 'manual-report']:
        if (OUT / ('before-' + name + '.json')).exists():
            check(name + '_entire_baseline_unchanged', read('before-' + name + '.json') == read('after-' + name + '.json'))
    sources = read('source-manifest.json')
    check('built_source_manifest_unchanged', all(hashlib.sha256((ROOT / name).read_bytes()).hexdigest() == expected for name, expected in sources.items()))
    for item in read('browser-downloads.json'):
        file = OUT / item['file']
        check(item['phase'] + '_actual_browser_bytes', file.read_bytes() == Path(item['original']).read_bytes())
        pdf = PdfReader(file)
        text = '\n'.join(page.extract_text() or '' for page in pdf.pages)
        (OUT / (item['phase'] + '-text.txt')).write_text(text)
        check(item['phase'] + '_expected_progress_text', item['expected_progress'] in text)
        check(item['phase'] + '_expected_weather_text', item['expected_weather'] in text)
        if item['expected_progress'] == 'Chưa ghi tiến độ':
            check(item['phase'] + '_no_false_zero_progress', '0%' not in text)
        check(item['phase'] + '_revision_number', re.search(r'Lần gửi\s+Gửi lúc\s+0*' + str(item['revision_no']) + r'\b', text) is not None)
        payload = read('after-' + ('manual-report' if item.get('source_kind') == 'manual' else 'report') + '.json')
        report = payload.get('data', payload)
        revision = next(row for row in report['revisions'] if row['revision_no'] == item['revision_no'])
        check(item['phase'] + '_exact_revision_hash_in_pdf', revision['snapshot_sha256'] in text)
        if item['phase'] in ['workflow1', 'workflow3']:
            old_file = WEATHER / ('BM02-lan-gui-' + str(item['revision_no']) + '-tieng-Viet.pdf')
            old_text = '\n'.join(page.extract_text() or '' for page in PdfReader(old_file).pages)
            check(item['phase'] + '_prior_pdf_text_only_missing_progress_label_changed', ' '.join(old_text.replace('0%', 'Chưa ghi tiến độ', 1).split()) == ' '.join(text.split()))
        result['files'][item['phase']] = {'file': item['file'], 'sha256': hashlib.sha256(file.read_bytes()).hexdigest(), 'bytes': file.stat().st_size, 'pages': len(pdf.pages), 'expected_progress': item['expected_progress']}
    result['bm02_parity'] = {}
    for item in read('side-downloads.json'):
        file = OUT / item['file']
        old_file = WEATHER / 'BM02-lan-gui-3-tieng-Viet.xlsx' if item['phase'] == 'workflow3' else ROOT / 'docs/acceptance-evidence/bm02-delivery-revision-development-2026-10-12/BM02-lan-2.xlsx'
        check(item['phase'] + '_bm02_actual_browser_bytes', file.read_bytes() == Path(item['original']).read_bytes())
        current, old = load_workbook(file), load_workbook(old_file)
        def cells(book):
            return {(sheet.title, cell.coordinate): (cell.value, cell.style_id, cell.number_format) for sheet in book for row in sheet for cell in row}
        check(item['phase'] + '_bm02_all_values_styles_unchanged', cells(current) == cells(old))
        check(item['phase'] + '_bm02_sheet_names_merges_unchanged', current.sheetnames == old.sheetnames and all(str(a.merged_cells) == str(b.merged_cells) for a, b in zip(current, old)))
        with ZipFile(file) as a, ZipFile(old_file) as b:
            def media(archive):
                return {name: archive.read(name) for name in archive.namelist() if name.startswith('xl/media/')}
            check(item['phase'] + '_bm02_embedded_media_unchanged', media(a) == media(b))
        result['bm02_parity'][item['phase']] = {'file': item['file'], 'progress_D10': current.worksheets[0]['D10'].value, 'weather_B6': current.worksheets[0]['B6'].value, 'quantity_G26': current.worksheets[0]['G26'].value}
        current.close()
        old.close()
    result['test_pdf_scope'] = {}
    for kind, label in [('workflow_absent', 'Chưa ghi tiến độ'), ('workflow_zero', '0%'), ('workflow_decimal', '42.5%'), ('manual_zero', '0%'), ('workflow_wrapper_boundary', '0%'), ('workflow_wrapper_overflow', 'Chưa ghi tiến độ')]:
        file = OUT / ('test-progress-' + kind + '.pdf')
        text = '\n'.join(page.extract_text() or '' for page in PdfReader(file).pages)
        check(kind + '_generated_test_pdf_text', label in text)
        result['test_pdf_scope'][kind] = {'file': file.name, 'sha256': hashlib.sha256(file.read_bytes()).hexdigest(), 'label': label, 'scope': 'generated pure-test PDF; not browser download or Development record'}
    result['images'] = {file.name: hashlib.sha256(file.read_bytes()).hexdigest() for file in sorted(OUT.iterdir()) if file.suffix in ['.jpg', '.png']}
    result['source'] = read('source-state.json')
    save('acceptance-checks.json', result)
    print(json.dumps({'checks': len(result['checks']), 'images': len(result['images']), 'files': result['files']}, ensure_ascii=False))


if __name__ == '__main__':
    mode = sys.argv[1]
    if mode == 'manifest':
        manifest()
    elif mode == 'verify':
        verify()
    else:
        baseline(mode)
