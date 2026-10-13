"""Verify retained actual BM04 downloads and native Google readbacks; no domain writes."""
import hashlib
import json
from pathlib import Path
import zipfile

import openpyxl

OUT = Path(__file__).resolve().parent
ROOT = OUT.parents[2]


def load(name):
    return json.loads((OUT / name).read_text())


def verify():
    before = {**load('before-database.json'), **load('before-construction-history.json')}
    after = {**load('after-database.json'), **load('after-construction-history.json')}
    changes, lost, additions = [], [], []
    count = 0
    for table, rows in before.items():
        if not isinstance(rows, list):
            continue
        old = {row['id']: row['md5'] for row in rows}
        current = {row['id']: row['md5'] for row in after[table]}
        count += len(old)
        lost.extend((table, key) for key in old if key not in current)
        changes.extend((table, key) for key in old if key in current and old[key] != current[key])
        additions.extend((table, key) for key in current if key not in old)
    source = load('source-manifest.json')['files']
    drift = [name for name, digest in source.items() if hashlib.sha256((ROOT / name).read_bytes()).hexdigest() != digest]
    workspace = load('before-workspace.json')
    action = workspace['actions'][0]
    check = next(row for row in workspace['checks'] if row['id'] == action['source_check_id'])
    expected_summary = 'TỔNG HỢP: Tổng lỗi: 1 | Đã khắc phục: 0 | Đang xử lý: 1'
    assertions = {
        'one_saved_open_uncoded_action': len(workspace['actions']) == 1 and action['status'] == 'open' and action['fault_code'] is None,
        'action_location_unset_linked_source_location_exists': action['location'] is None and check['location'] == '[Synthetic] Khu B — ô kiểm tra 04',
        'no_existing_row_changed_or_lost': not changes and not lost,
        'workspace_source_records_unchanged': workspace['actions'] == load('after-workspace.json')['actions'] and workspace['checks'] == load('after-workspace.json')['checks'],
        'no_source_change': not drift,
        'migration_count_unchanged': before['migration_count'] == after['migration_count'] == 158,
    }
    original = openpyxl.load_workbook('/Users/khiemnguyen/Works/ttcons/documents/Bo_Bieu_Mau_Giam_Sat_Thu_Thiem.xlsx', read_only=False)
    headers = [original['BM-GS-04'].cell(7, c).value for c in range(1, 11)]
    inspected = load('download-inspection.json')
    mismatches, errors, populated, cells_checked = [], [], 0, 0
    files = {}
    all_local = {}
    for label in ['list', 'detail']:
        file = OUT / ('BM04-' + label + '.xlsx')
        workbook = openpyxl.load_workbook(file, read_only=False, data_only=False)
        sheet = workbook.worksheets[0]
        downloaded = Path(inspected[label]['download_path'])
        assertions[label + '_actual_download_byte_equal'] = file.read_bytes() == downloaded.read_bytes()
        assertions[label + '_client_headers_preserved'] = [sheet.cell(7, c).value for c in range(1, 11)] == headers
        assertions[label + '_location_matches_check'] = sheet['C8'].value == check['location']
        assertions[label + '_missing_code_stays_blank_numeric_stt_counts'] = sheet['B8'].value is None and sheet['A8'].value == 1 and 'COUNT(A8:A10)' in sheet['A13'].value
        assertions[label + '_status_matches_native_source'] = sheet['I8'].value == 'Mới' and action['status'] == 'open'
        assertions[label + '_three_signature_placeholders_retained'] = all(sheet.cell(18, c).value == '(Ký & ghi rõ họ tên)' for c in [1, 4, 7])
        cells = {}
        for suffix in ['cells', 'footer']:
            response = load('google-' + label + '-' + suffix + '.json')['structuredContent']
            for tab in response['sheets']:
                for data in tab.get('data', []):
                    for i, row in enumerate(data.get('rowData', []), start=data.get('startRow', 0) + 1):
                        for j, cell in enumerate(row.get('values', []), start=data.get('startColumn', 0) + 1):
                            cells[(tab['properties']['title'], i, j)] = cell
        assertions[label + '_google_summary_recalculated'] = cells[(sheet.title, 13, 1)]['effectiveValue']['stringValue'] == expected_summary
        local = {}
        for tab_name, max_row, max_col in [(sheet.title, 18, 10), ('Hình ảnh', 25, 8)]:
            tab = workbook[tab_name]
            for i in range(1, max_row + 1):
                for j in range(1, max_col + 1):
                    cells_checked += 1
                    key = (tab_name, i, j)
                    value = tab.cell(i, j).value
                    value = None if value in [None, ''] else value
                    local[key] = value
                    populated += value is not None
                    cell = cells.get(key, {})
                    effective = cell.get('effectiveValue', {})
                    actual = effective.get('stringValue', effective.get('numberValue', effective.get('boolValue')))
                    if isinstance(value, str) and value.startswith('='):
                        actual = cell.get('userEnteredValue', {}).get('formulaValue')
                    if actual != value:
                        mismatches.append([label, tab_name, i, j, value, actual])
                    if 'errorValue' in effective:
                        errors.append([label, tab_name, i, j, effective['errorValue']])
        all_local[label] = local
        with zipfile.ZipFile(file) as package:
            media = {name: hashlib.sha256(package.read(name)).hexdigest() for name in package.namelist() if name.startswith('xl/media/')}
        files[label] = {'sha256': hashlib.sha256(file.read_bytes()).hexdigest(), 'media': media}
    assertions['both_exports_have_identical_cells_and_formulas'] = all_local['list'] == all_local['detail']
    assertions['same_embedded_source_photo_in_both_exports'] = files['list']['media'] == files['detail']['media'] and len(files['list']['media']) == 1
    assertions['google_all_read_values_and_formulas_match'] = not mismatches
    assertions['google_no_effective_errors'] = not errors
    return {'assertions': assertions, 'passed': all(assertions.values()), 'old_unique_rows_checked': count, 'old_row_changes': changes, 'old_rows_lost': lost, 'new_rows': additions, 'source_files_checked': len(source), 'source_drift': drift, 'google_cells_checked': cells_checked, 'google_populated_cells': populated, 'google_mismatches': mismatches, 'google_errors': errors, 'files': files, 'ui_findings': [load('browser-list-layout-measurement.json')], 'scope': 'Export technical checks PASS for Admin synthetic Development: one existing open action with no fault code and inherited linked-check location; actual list/detail downloads and native Google import. List UI P2 overlap at1280 unresolved; NOT UX ACCEPTED. Other lifecycle rows, ordinary roles, capacity, native Excel and print NOT_RUN.'}


if __name__ == '__main__':
    result = verify()
    (OUT / 'verification.json').write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n')
    print(json.dumps(result, ensure_ascii=False, indent=2))
    assert result['passed']
