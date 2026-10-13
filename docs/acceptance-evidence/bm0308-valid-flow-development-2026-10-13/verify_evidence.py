"""Verify the source-to-download-to-native-Google proof, without workbook edits."""
import hashlib
import json
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter

OUT = Path(__file__).resolve().parent
ROOT = OUT.parents[2]
checks = []


def check(label, condition):
    checks.append({'label': label, 'passed': bool(condition)})
    assert condition, label


def read(name):
    return json.loads((OUT / name).read_text())


before, after = read('before-database.json'), read('after-database.json')
check('database and migration count unchanged', before['database'] == after['database'] and before['migration_count'] == after['migration_count'] == 157)
for table, rows in before.items():
    if not isinstance(rows, list): continue
    current = {r['id']: r['row_md5'] for r in after[table]}
    check('retained rows unchanged: ' + table, all(current.get(r['id']) == r['row_md5'] for r in rows))

manifest = read('source-manifest.json')
check('968 application files unchanged', len(manifest) == 968 and all(hashlib.sha256((ROOT / p).read_bytes()).hexdigest() == h for p, h in manifest.items()))

state = read('fixture-state.json')
handover = read('handover-completed-detail.json')
acceptance = read('final-acceptance-final-detail.json')
form = read('final-source-signed-pdf.json')
check('final acceptance completed pass with derived actual signature/PDF links', acceptance['status'] == 'completed' and acceptance['result'] == 'pass' and acceptance['signature_status'] == 'signed' and acceptance['document_status'] == 'linked' and acceptance['form_instance_id'] == form['instance']['id'] and acceptance['pdf_export_id'] == form['pdf_exports'][0]['id'])
check('internal test signature belongs to same submitted form/revision', len(form['signatures']) == 1 and form['signatures'][0]['purpose'] == 'internal_ack' and form['instance']['status'] == 'submitted')
check('handover actually completed, correct project and final acceptance', handover['status'] == 'handed_over' and handover['completed_at'] and handover['final_acceptance_id'] == acceptance['id'] and handover['project_id'] == acceptance['project_id'] == state['handover']['project_id'])
snapshot = handover['readiness_snapshot_json']
check('six verified required items; no readiness blockers', snapshot['ready'] and snapshot['blockers'] == [] and len(snapshot['checklist_items']) == 6 and all(i['status'] == 'verified' and i['required'] and i['physical_result'] == 'pass' for i in snapshot['checklist_items']))
check('warehouse/documents ready from source checks', handover['warehouse_closed'] and handover['documents_ready'] and all(snapshot['counts'][k] == 0 for k in ['open_issues', 'open_ncrs', 'open_reworks', 'open_construction_actions', 'pending_surplus']))
for i in snapshot['checklist_items']:
    if i['item_code'] in state['handover']['documents']:
        doc = read(i['item_code'] + '-uploaded.json')
        check('manual item linked to actual same-project document: ' + i['item_code'], i['document_id'] == doc['id'] and doc['project_id'] == handover['project_id'] and doc['status'] == 'active')
check('client package own reconciliation unsupported is explicit', read('client-reconciliation-after-form.json')['available'] is False)
check('builtin reconciliation submitted against exact own source', read('builtin-reconciliation-submitted.json')['reconciliation']['status'] == 'submitted' and read('builtin-reconciliation-submitted.json')['reconciliation']['form_instance_id'] == state['builtin']['form_id'])
official = read('official-client-work-item-readonly.json')
check('official source readiness captured read-only, no completion claim', official['pipeline_engine_version'] == 4 and official['business_source']['ready'] is False and official['status'] == 'pending')
run = read('handover-final-pipeline-run.json')
check('isolated template is engine2; no official gate parity claim', run['engine_version'] == 2)
wi = read('handover-final-project-work-items.json')
check('accounting closure intentionally not completed', next(i for i in wi if i['title'] == 'Xác nhận bàn giao hồ sơ kế toán')['status'] != 'completed')

downloads = [
    ('bm03builtin', 'BM03-builtin-submitted.xlsx', '/Users/khiemnguyen/Downloads/BM-GS-03 (10).xlsx'),
    ('bm03client', 'BM03-client-submitted.xlsx', '/Users/khiemnguyen/Downloads/BM-GS-03 (11).xlsx'),
    ('bm08', 'BM08-handed-over.xlsx', '/Users/khiemnguyen/Downloads/BM-GS-08 (7).xlsx')]
files = []
for key, name, downloaded in downloads:
    data = (OUT / name).read_bytes()
    check('byte-equal actual browser download: ' + key, data == Path(downloaded).read_bytes())
    w = load_workbook(OUT / name, data_only=False)
    check('one native source sheet: ' + key, len(w.worksheets) == 1)
    sheet = w.worksheets[0]
    google = read(key + '-google-cells-full.json')['structuredContent']
    gs = google['sheets'][0]
    check('exact source sheet imported: ' + key, gs['properties']['title'] == sheet.title)
    observed = {}
    errors = []
    for block in gs['data']:
        for ri, row in enumerate(block.get('rowData', []), block.get('startRow', 0) + 1):
            for ci, cell in enumerate(row.get('values', []), block.get('startColumn', 0) + 1):
                coord = get_column_letter(ci) + str(ri)
                ev = cell.get('effectiveValue', {})
                if 'errorValue' in ev: errors.append([coord, ev['errorValue']])
                observed[coord] = cell
    check('zero effective cell errors including VALUE: ' + key, not errors)
    count = 0
    for row in sheet:
        for cell in row:
            if cell.value is None or cell.value == '': continue
            got = observed.get(cell.coordinate, {}).get('effectiveValue', {})
            if isinstance(cell.value, (int, float)):
                check(key + ' numeric source value ' + cell.coordinate, got.get('numberValue') == cell.value)
            else:
                check(key + ' text source value ' + cell.coordinate, got.get('stringValue') == cell.value)
            count += 1
    check('all60 populated cells verified: ' + key, count == 60)
    if key != 'bm08':
        answers = read(('builtin' if key == 'bm03builtin' else 'client') + '-form-answers.json')
        for ri, item in enumerate(answers['checklist_items'], 9):
            check(key + ' separates measured versus reconciled quantity row' + str(ri), sheet['E' + str(ri)].value == item['design_quantity'] and sheet['F' + str(ri)].value == item['actual_quantity'] and item['actual_quantity'] != item['confirmed_quantity'])
            check(key + ' criteria and explicit quality source row' + str(ri), item['standard_reference'] in sheet['C' + str(ri)].value and item['drawing_reference'] in sheet['C' + str(ri)].value and sheet['D' + str(ri)].value == item['technical_requirement'] and item['quality_note'] in sheet['G' + str(ri)].value)
        check(key + ' conclusion matches source', '[ X ] ĐỒNG Ý NGHIỆM THU' in sheet['A13'].value)
    else:
        check('BM08 six native rows and source conclusion', [sheet['D' + str(i)].value for i in range(8,14)] == ['Đạt'] * 6 and '[ X ] ĐỦ ĐIỀU KIỆN BÀN GIAO' in sheet['A15'].value)
        check('BM08 no invented defect deadline', all(not sheet['F' + str(i)].value for i in range(8,14)))
    files.append({'name': name, 'downloaded_path': downloaded, 'sha256': hashlib.sha256(data).hexdigest(), 'google_id': google['spreadsheetId'], 'source_range': sheet.calculate_dimension(), 'verified_populated_cells': count, 'effective_errors': errors})

result = {'checks_passed': len(checks), 'checks': checks, 'files': files, 'scope': 'Isolated synthetic Development dossier/export UAT. Full official Rev10 and accounting closure NOT_RUN. User visual acceptance pending. No workbook mutation.'}
(OUT / 'evidence-verification.json').write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n')
print(json.dumps({'checks_passed': len(checks), 'files': files}, ensure_ascii=False))
