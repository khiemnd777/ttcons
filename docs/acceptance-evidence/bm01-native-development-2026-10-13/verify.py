"""Read-only audit of the actual native BM01 download, Google readback and history."""
import hashlib
import json
from pathlib import Path
import openpyxl

OUT = Path(__file__).resolve().parent
ROOT = OUT.parents[2]


def load(name):
    return json.loads((OUT / name).read_text())


def verify():
    source = load('source-form.json')
    answers = source['answers_snapshot_json']
    row = answers['dimension_items'][0]
    project = load('source-project.json')
    workbook = openpyxl.load_workbook(OUT / 'BM01.xlsx', data_only=False)
    sheet = workbook['2026-10-09_01']
    client = load('client-template.json')['BM-GS-01']
    downloaded = load('download.json')
    imported = load('google-import.json')['structuredContent']
    metadata = load('google-metadata.json')['structuredContent']
    google = load('google-cells.json')['structuredContent']
    cells = {}
    errors = []
    for tab in google['sheets']:
        for block in tab['data']:
            for ri, rd in enumerate(block.get('rowData', []), block.get('startRow', 0) + 1):
                for ci, cell in enumerate(rd.get('values', []), block.get('startColumn', 0) + 1):
                    entered = cell.get('userEnteredValue', {})
                    effective = cell.get('effectiveValue', {})
                    value = entered.get('formulaValue', entered.get('stringValue', entered.get('numberValue', entered.get('boolValue'))))
                    cells[(ri, ci)] = value
                    if 'errorValue' in effective or cell.get('formattedValue', '').startswith(('#VALUE!', '#REF!', '#DIV/0!', '#N/A', '#NAME?', '#NUM!', '#ERROR!')):
                        errors.append((ri, ci, cell))
    parity = all(cells.get((r, c)) == sheet.cell(r, c).value for r in range(1, 21) for c in range(1, 8))
    before = {**load('before-database.json'), **load('before-construction-history.json')}
    after = {**load('after-database.json'), **load('after-construction-history.json')}
    changes, additions, old_count = [], [], 0
    for table, records in before.items():
        if not isinstance(records, list):
            continue
        old = {r['id']: r['md5'] for r in records}
        now = {r['id']: r['md5'] for r in after[table]}
        old_count += len(old)
        changes.extend((table, k) for k in old if now.get(k) != old[k])
        additions.extend((table, k) for k in now if k not in old)
    source_files = load('source-before.json')['files']
    current = {name: hashlib.sha256((ROOT / name).read_bytes()).hexdigest() for name in source_files}
    audit = load('new-audit-rows.json')
    expected = [1, row['area'], row['dimension'] + ' • ' + row['unit'], row['mep'], row['existing_condition'], row['assessment'], row['note']]
    drive = load('google-drive-metadata.json')['structuredContent']
    assertions = {
        'actual_button_download_retained_byte_for_byte': (OUT / 'BM01.xlsx').read_bytes() == Path(downloaded['sourcePath']).read_bytes(),
        'exact_linked_source_form_work_item_project': source['work_item_id'] == load('source-work-item.json')['id'] and source['project_id'] == project['id'],
        'exact_template_and_version_three_pin': source['form_template_id'] == load('source-template.json')['id'] and source['form_version_id'] == load('source-version.json')['id'] and load('source-version.json')['version_no'] == 3,
        'draft_status_and_revision_zero_preserved': source['status'] == 'draft' and source['current_revision'] == 0,
        'one_saved_dimension_row_and_one_exported_row': len(answers['dimension_items']) == 1 and [sheet.cell(10, c).value for c in range(1, 8)] == expected,
        'project_name_code_address_match_source': sheet['B5'].value == project['name'] and sheet['F5'].value == project['code'] and sheet['B6'].value == project['address'],
        'measurement_date_uses_saved_business_date': sheet['F6'].value == answers['measurement_date'] == '2026-10-09',
        'survey_unit_and_client_representative_match_source': sheet['B7'].value == answers['survey_unit'] and sheet['F7'].value == answers['client_representative'],
        'all_seven_original_client_headers_preserved': all(sheet.cell(9, c).value == client[sheet.cell(9, c).coordinate] for c in range(1, 8)),
        'original_title_and_signature_placeholders_preserved': all(sheet[a].value == client[a] for a in ['A1', 'A2', 'A3', 'A17', 'C17', 'E17', 'A18', 'C18', 'E18']),
        'all_original_sample_rows_cleared': all(sheet.cell(r, c).value is None for r in range(11, 15) for c in range(1, 8)),
        'single_output_tab_no_added_photo_columns_or_sheet': workbook.sheetnames == ['2026-10-09_01'] and sheet.max_column == 7 and not sheet._images,
        'native_google_conversion_confirmed': imported['converted'] and imported['mimeType'] == 'application/vnd.google-apps.spreadsheet',
        'google_tab_identity_matches_download': len(metadata['sheets']) == 1 and metadata['sheets'][0]['properties']['title'] == sheet.title,
        'all_140_bounded_google_cells_equal_xlsx': parity,
        'no_google_effective_error_or_value_error': not errors,
        'new_google_file_private': drive['shared'] is False and all(p['type'] == 'user' and p['role'] == 'owner' for p in drive['permissions']),
        'all_baseline_old_rows_unchanged_and_retained': not changes,
        'only_one_native_source_export_audit_added': len(additions) == 1 and additions[0][0] == 'audit_logs' and len(audit) == 1 and audit[0]['event_type'] == 'native_source_exported',
        'migration_158_and_database_unchanged': before['migration_count'] == after['migration_count'] == 158 and before['database'] == after['database'] == 'consflow',
        'all_storage_files_preserved': load('storage-before.json') == load('storage-after.json'),
        'all_978_built_source_hashes_unchanged': len(current) == 978 and current == source_files,
    }
    return {'passed': all(assertions.values()), 'assertions': assertions, 'old_unique_rows_checked': old_count, 'changed_or_lost_rows': changes, 'new_rows': additions, 'bounded_google_cells_compared': 140, 'populated_google_cells': sum(v is not None for v in cells.values()), 'google_errors': errors, 'source_files_checked': len(current), 'storage_files_checked': len(load('storage-before.json')), 'formula_count': sum(c.data_type == 'f' for row in sheet for c in row), 'limits': 'Development Admin; one saved native survey v3 draftR0, one measurement row, no evidence. Browser desktop and unchanged native Google import. Not submitted/Chot ho so/full Rev10 completion. Alternate client_site_measurement_record, multirow/capacity, mobile/physical device, other roles, Excel desktop/print NOT_RUN. User visual sign-off pending.'}


if __name__ == '__main__':
    result = verify()
    (OUT / 'verification.json').write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n')
    print(json.dumps(result, ensure_ascii=False, indent=2))
    assert result['passed']
