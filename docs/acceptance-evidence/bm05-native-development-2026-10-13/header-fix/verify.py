"""Audit the approved BM05 header repair without changing source/workbooks/Google."""
import hashlib
import json
import posixpath
from pathlib import Path
import xml.etree.ElementTree as ET
import zipfile
import openpyxl

OUT = Path(__file__).resolve().parent
PRIOR = OUT.parent
ROOT = OUT.parents[3]


def load(name, base=OUT):
    return json.loads((base / name).read_text())


def column_geometry(path):
    with zipfile.ZipFile(path) as package:
        workbook = ET.fromstring(package.read('xl/workbook.xml'))
        workbook_ns = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
        rid = workbook.find(workbook_ns + 'sheets')[0].attrib['{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id']
        relationships = ET.fromstring(package.read('xl/_rels/workbook.xml.rels'))
        target = next(r.attrib['Target'] for r in relationships if r.attrib['Id'] == rid)
        part = target.lstrip('/') if target.startswith('/') else posixpath.normpath(posixpath.join('xl', target))
        assert part.startswith('xl/worksheets/')
        sheet = ET.fromstring(package.read(part))
    ns = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
    return [c.attrib for c in sheet.find(ns + 'cols')]


def verify():
    current_path = OUT / 'BM05.xlsx'
    old_path = PRIOR / 'BM05.xlsx'
    current = openpyxl.load_workbook(current_path).active
    old = openpyxl.load_workbook(old_path).active
    values = lambda w: {c.coordinate: c.value for row in w for c in row if c.value is not None}
    cells, errors = {}, []
    google = load('google-cells.json')['structuredContent']
    for tab in google['sheets']:
        for block in tab['data']:
            for r, row in enumerate(block.get('rowData', []), block.get('startRow', 0) + 1):
                for c, cell in enumerate(row.get('values', []), block.get('startColumn', 0) + 1):
                    entered = cell.get('userEnteredValue', {})
                    cells[(r, c)] = entered.get('formulaValue', entered.get('stringValue', entered.get('numberValue', entered.get('boolValue'))))
                    if 'errorValue' in cell.get('effectiveValue', {}) or cell.get('formattedValue', '').startswith(('#VALUE!', '#REF!', '#DIV/0!', '#N/A', '#NAME?', '#NUM!', '#ERROR!')):
                        errors.append((r, c, cell))
    before = {**load('before-database.json'), **load('before-construction-history.json')}
    after = {**load('after-database.json'), **load('after-construction-history.json')}
    changes, additions, checked = [], [], 0
    for table, rows in before.items():
        if not isinstance(rows, list):
            continue
        original = {r['id']: r['md5'] for r in rows}
        now = {r['id']: r['md5'] for r in after[table]}
        checked += len(original)
        changes.extend((table, key) for key in original if now.get(key) != original[key])
        additions.extend((table, key) for key in now if key not in original)
    built = load('source-build-manifest.json')['files']
    actual = {name: hashlib.sha256((ROOT / name).read_bytes()).hexdigest() for name in built}
    originals = load('source-before.json')['files']
    delta = [n for n in originals if actual[n] != originals[n]]
    source = load('source-form.json', PRIOR)
    risk = source['answers_snapshot_json']['risk_items'][0]
    expected = [1, risk['issue_code'], risk['detected_date'], risk['risk'], risk['cause'], risk['impact'], 'Trung bình', risk['mitigation'], risk['owner'], risk['due_date'], risk['status']]
    client = load('client-template.json', PRIOR)['BM-GS-05']
    drive = load('google-drive-metadata.json')['structuredContent']
    imported = load('google-import.json')['structuredContent']
    metadata = load('google-metadata.json')['structuredContent']
    runtime = load('runtime.json')
    audit = load('new-audit-rows.json')
    old_merges = {str(m) for m in old.merged_cells}
    current_merges = {str(m) for m in current.merged_cells}
    assertions = {
        'actual_browser_download_retained_byte_for_byte': current_path.read_bytes() == Path(load('download.json')['sourcePath']).read_bytes(),
        'original_pre_repair_file_preserved': hashlib.sha256(old_path.read_bytes()).hexdigest() == load('workbook-inspection.json', PRIOR)['sha256'],
        'only_added_merge_is_b5_k5': current_merges == old_merges | {'B5:K5'},
        'all_cell_values_equal_pre_repair': values(current) == values(old),
        'full_project_name_retained': current['B5'].value == old['B5'].value and '[DEV-BM-REUSE-20261009]' in current['B5'].value,
        'row5_height_reduced_for_actual_fixture': current.row_dimensions[5].height == 33 and old.row_dimensions[5].height > 169,
        'all_risk_column_geometry_preserved': column_geometry(current_path) == column_geometry(old_path),
        'all_eleven_client_headers_preserved': all(current.cell(7, c).value == client[current.cell(7, c).coordinate] for c in range(1, 12)),
        'saved_risk_values_dates_and_status_preserved': [current.cell(8, c).value for c in range(1, 12)] == expected,
        'all_original_signature_placeholders_preserved': all(current[a].value == client[a] for a in ['A12', 'D12', 'G12', 'A13', 'D13', 'G13']),
        'exact_native_form_pin_answers_status_and_revision_unchanged': load('source-form-after.json') == source,
        'native_private_google_conversion_confirmed': imported['converted'] and imported['mimeType'] == 'application/vnd.google-apps.spreadsheet' and drive['shared'] is False and all(p['type'] == 'user' and p['role'] == 'owner' for p in drive['permissions']),
        'single_google_tab_matches_actual_download': len(metadata['sheets']) == 1 and metadata['sheets'][0]['properties']['title'] == current.title,
        'all_165_google_cells_match_xlsx': all(cells.get((r, c)) == current.cell(r, c).value for r in range(1, 16) for c in range(1, 12)),
        'no_google_effective_error_or_value_error': not errors,
        'google_ui_confirms_b5_k5_and_full_name': 'Value: B5:K5' in (OUT / 'google-merged-cell-state.txt').read_text() and current['B5'].value in (OUT / 'google-merged-cell-state.txt').read_text(),
        'all_2893_old_unique_rows_unchanged_and_retained': not changes and checked == 2893,
        'only_one_native_source_export_audit_added': len(additions) == 1 and additions[0][0] == 'audit_logs' and len(audit) == 1 and audit[0]['event_type'] == 'native_source_exported',
        'database_consflow_migration_158_unchanged': before['database'] == after['database'] == 'consflow' and before['migration_count'] == after['migration_count'] == 158,
        'all_149_storage_files_unchanged': load('storage-before.json') == load('storage-after.json') and len(load('storage-before.json')) == 149,
        'all_979_current_source_hashes_match_built_manifest': actual == built and len(actual) == 979,
        'only_approved_existing_renderer_changed': delta == ['backend/internal/core/bm_gs_workbook.go'] and set(actual) - set(originals) == {'backend/internal/core/bm_gs_risk_header_test.go'},
        'only_backend_recreated_with_changed_image': runtime['only_backend_recreated'] and runtime['backend_image_changed'],
        'development_build_and_readiness_passed': runtime['build_exit_code'] == runtime['update_exit_code'] == 0 and runtime['readiness_http'] == runtime['frontend_http'] == 200,
    }
    return {'passed': all(assertions.values()), 'assertions': assertions, 'old_unique_rows_checked': checked, 'changed_or_lost_rows': changes, 'new_rows': additions, 'source_files_checked': len(actual), 'existing_source_delta': delta, 'bounded_google_cells_compared': 165, 'populated_google_cells': sum(v is not None for v in cells.values()), 'google_errors': errors, 'before_row5_height': old.row_dimensions[5].height, 'after_row5_height': current.row_dimensions[5].height, 'limits': 'Development Admin, same native risk v4 draftR0 with one saved risk row, actual desktop download and unchanged private nativeGoogle. Merged-cell UI and100%layout screenshots; user review pending. Submitted/resolved/fullRev10/contextual/multirow/capacity/mobile/otherroles/Excel/printNOT_RUN. Source tests cover short/long frozen names and other-template geometry without runtime data mutation.'}


if __name__ == '__main__':
    result = verify()
    (OUT / 'verification.json').write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n')
    print(json.dumps(result, ensure_ascii=False, indent=2))
    assert result['passed']
