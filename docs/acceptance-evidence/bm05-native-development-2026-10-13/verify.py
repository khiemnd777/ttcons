"""Read-only BM05 source/download/Google/history audit; layout review is separate."""
import hashlib
import json
from pathlib import Path
import openpyxl

OUT = Path(__file__).resolve().parent
ROOT = OUT.parents[2]


def load(name):
    return json.loads((OUT / name).read_text())


def verify():
    source = load('source-form.json')
    answers = source['answers_snapshot_json']
    risk = answers['risk_items'][0]
    project = load('source-project.json')
    wi = load('source-work-item.json')
    workbook = openpyxl.load_workbook(OUT / 'BM05.xlsx', data_only=False)
    sheet = workbook.active
    client = load('client-template.json')['BM-GS-05']
    fields = load('source-fields.json')[0]['validation_json']['item_fields']
    severity_labels = {o['value']: o['label'] for f in fields if f['key'] == 'severity' for o in f['options']}
    download = load('download.json')
    imported = load('google-import.json')['structuredContent']
    metadata = load('google-metadata.json')['structuredContent']
    google = load('google-cells.json')['structuredContent']
    drive = load('google-drive-metadata.json')['structuredContent']
    cells, errors = {}, []
    for tab in google['sheets']:
        for block in tab['data']:
            for ri, rd in enumerate(block.get('rowData', []), block.get('startRow', 0) + 1):
                for ci, cell in enumerate(rd.get('values', []), block.get('startColumn', 0) + 1):
                    entered = cell.get('userEnteredValue', {})
                    effective = cell.get('effectiveValue', {})
                    cells[(ri, ci)] = entered.get('formulaValue', entered.get('stringValue', entered.get('numberValue', entered.get('boolValue'))))
                    if 'errorValue' in effective or cell.get('formattedValue', '').startswith(('#VALUE!', '#REF!', '#DIV/0!', '#N/A', '#NAME?', '#NUM!', '#ERROR!')):
                        errors.append((ri, ci, cell))
    before = {**load('before-database.json'), **load('before-construction-history.json')}
    after = {**load('after-database.json'), **load('after-construction-history.json')}
    changes, additions, count = [], [], 0
    for table, records in before.items():
        if not isinstance(records, list):
            continue
        old = {r['id']: r['md5'] for r in records}
        now = {r['id']: r['md5'] for r in after[table]}
        count += len(old)
        changes.extend((table, k) for k in old if now.get(k) != old[k])
        additions.extend((table, k) for k in now if k not in old)
    hashes = load('source-before.json')['files']
    actual = {name: hashlib.sha256((ROOT / name).read_bytes()).hexdigest() for name in hashes}
    audit = load('new-audit-rows.json')
    expected = [1, risk['issue_code'], risk['detected_date'], risk['risk'], risk['cause'], risk['impact'], severity_labels[risk['severity']], risk['mitigation'], risk['owner'], risk['due_date'], risk['status']]
    assertions = {
        'actual_native_button_download_retained_byte_for_byte': (OUT / 'BM05.xlsx').read_bytes() == Path(download['sourcePath']).read_bytes(),
        'same_project_linked_form_work_item': source['work_item_id'] == wi['id'] and source['project_id'] == wi['project_id'] == project['id'],
        'exact_native_risk_template_and_version_four_pin': load('source-template.json')['code'] == 'client_site_risk_assessment' and source['form_template_id'] == load('source-template.json')['id'] and source['form_version_id'] == load('source-version.json')['id'] and load('source-version.json')['version_no'] == 4,
        'saved_draft_revision_zero_preserved': source['status'] == 'draft' and source['current_revision'] == 0,
        'one_saved_risk_row_equals_all_eleven_output_cells': len(answers['risk_items']) == 1 and [sheet.cell(8, c).value for c in range(1, 12)] == expected,
        'source_project_name_and_code_match': sheet['B5'].value == '[' + project['code'] + '] ' + project['name'],
        'detected_and_due_dates_are_saved_dates': sheet['C8'].value == risk['detected_date'] == '2026-10-09' and sheet['J8'].value == risk['due_date'] == '2026-10-10',
        'risk_severity_uses_correct_vietnamese_label': sheet['G8'].value == severity_labels[risk['severity']] == 'Trung bình',
        'status_is_saved_risk_text_not_draft_or_work_item_state': sheet['K8'].value == risk['status'] == 'Đang theo dõi synthetic' and sheet['K8'].value not in [source['status'], wi['status']],
        'all_eleven_original_client_headers_preserved': all(sheet.cell(7, c).value == client[sheet.cell(7, c).coordinate] for c in range(1, 12)),
        'original_title_and_three_signature_placeholders_preserved': all(sheet[a].value == client[a] for a in ['A1', 'A2', 'A3', 'A12', 'D12', 'G12', 'A13', 'D13', 'G13']),
        'original_second_sample_row_cleared': all(sheet.cell(9, c).value is None for c in range(1, 12)),
        'one_sheet_eleven_columns_no_photo_column_or_appendix': len(workbook.sheetnames) == 1 and sheet.max_column == 11 and not sheet._images,
        'native_google_conversion_confirmed': imported['converted'] and imported['mimeType'] == 'application/vnd.google-apps.spreadsheet',
        'native_google_tab_identity_matches_download': len(metadata['sheets']) == 1 and metadata['sheets'][0]['properties']['title'] == sheet.title,
        'all_165_bounded_google_cells_equal_xlsx': all(cells.get((r, c)) == sheet.cell(r, c).value for r in range(1, 16) for c in range(1, 12)),
        'no_google_error_value_or_value_error': not errors,
        'new_google_file_private_owner_only': drive['shared'] is False and all(p['type'] == 'user' and p['role'] == 'owner' for p in drive['permissions']),
        'all_old_business_workflow_construction_audit_rows_retained': not changes,
        'only_one_native_source_export_audit_added': len(additions) == 1 and additions[0][0] == 'audit_logs' and len(audit) == 1 and audit[0]['event_type'] == 'native_source_exported',
        'database_consflow_migration_158_unchanged': before['database'] == after['database'] == 'consflow' and before['migration_count'] == after['migration_count'] == 158,
        'all_149_storage_files_unchanged': load('storage-before.json') == load('storage-after.json') and len(load('storage-before.json')) == 149,
        'all_978_built_source_hashes_unchanged': actual == hashes and len(actual) == 978,
    }
    return {'passed': all(assertions.values()), 'assertions': assertions, 'old_unique_rows_checked': count, 'changed_or_lost_rows': changes, 'new_rows': additions, 'bounded_google_cells_compared': 165, 'populated_google_cells': sum(v is not None for v in cells.values()), 'google_errors': errors, 'formula_count': sum(c.data_type == 'f' for row in sheet for c in row), 'source_files_checked': len(actual), 'storage_files_checked': len(load('storage-before.json')), 'layout_observation': load('layout-observation.json'), 'acceptance_status': 'DATA TECHNICALLY VERIFIED; visual issue: project B5 wraps into many lines, row5 height169.35pt; NOT UX ACCEPTED.', 'limits': 'Development Admin, one saved native risk v4 draftR0, one risk row, no evidence. Actual desktop browser download and unchanged native Google import. Submitted/resolved/fullRev10/contextual sources/multirow/capacity/mobile/otherroles/Excel/print NOT_RUN. No application or workbook repair.'}


if __name__ == '__main__':
    result = verify()
    (OUT / 'verification.json').write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n')
    print(json.dumps(result, ensure_ascii=False, indent=2))
    assert result['passed']
