"""Read-only checks of canonical Development records and actual browser downloads.

Run with the bundled Python runtime and the temporary API helper/private session.
The script writes evidence JSON. It never edits an XLSX, Sheet or application record.
"""
import hashlib
import importlib.util
import io
import json
from pathlib import Path
import subprocess
import urllib.request
from zipfile import ZipFile
import xml.etree.ElementTree as ET

from openpyxl import load_workbook
from PIL import Image

spec = importlib.util.spec_from_file_location('acceptance_api', '/private/tmp/consflow-bm02-delivery-acceptance.py')
api = importlib.util.module_from_spec(spec)
spec.loader.exec_module(api)
out, root = api.OUT, api.ROOT
state = json.loads((out / 'fixture-state.json').read_text())
result = {'scope': 'Canonical Development manual-report delivery revisions; no application or schema edits', 'checks': {}}

def check(name, value):
    result['checks'][name] = bool(value)
    assert value, name

before = json.loads((out / 'delivery1-before.json').read_text())
after = api.call('/work-reports/daily-reports/' + state['source_id'])
api.save('delivery2-after.json', after)
r1 = next(r for r in after['revisions'] if r['revision_no'] == 1)
r2 = next(r for r in after['revisions'] if r['revision_no'] == 2)
check('same_report_two_revisions', after['work_report_id'] == state['work_report_id'] and after['current_revision_no'] == 2 and len(after['revisions']) == 2)
check('revision1_snapshot_hash_unchanged', r1['snapshot_sha256'] == before['revisions'][0]['snapshot_sha256'])
check('revision1_snapshot_json_unchanged', r1['snapshot_json'] == before['revisions'][0]['snapshot_json'])
check('revision1_entire_revision_record_unchanged', r1 == before['revisions'][0])
check('revision1_photo_membership', [p['id'] for p in r1['evidence_records']] == [state['photo1_id']])
check('revision2_photo_membership', [p['id'] for p in r2['evidence_records']] == [state['photo2_id']])
check('recipient_review_actor', any(e['event_type'] == 'needs_changes' and e['actor_id'] == state['recipient_id'] for e in after['events']))
check('revision2_distinct_snapshot_hash', r2['snapshot_sha256'] != r1['snapshot_sha256'])
result['snapshot_sha256'] = {'revision1': r1['snapshot_sha256'], 'revision2': r2['snapshot_sha256']}

workspace = api.call('/projects/' + state['project_id'] + '/construction-workspace')
api.save('workspace-after.json', workspace)
workspace_before = json.loads((out / 'workspace-baseline.json').read_text())
check('workflow_runtime_unchanged', workspace['runtime'] == workspace_before['runtime'])
check('work_items_unchanged', workspace['work_items'] == workspace_before['work_items'])
db_after = api.baseline_db()
api.save('database-after.json', db_after)
db_before = json.loads((out / 'database-before.json').read_text())
check('database_projects_forms_migrations_preserved', all(db_after[k] == db_before[k] for k in ['database', 'projects', 'forms', 'migration_count']))
after_hashes = {r['id']: r for r in db_after['preexisting_revision_hashes'] or []}
if db_before['preexisting_revision_hashes']:
    check('preexisting_delivery_revisions_unchanged', all(after_hashes.get(r['id']) == r for r in db_before['preexisting_revision_hashes']))
else:
    result['not_applicable'] = ['No delivery revision existed before this fixture; prior-revision preservation check has an empty baseline. Revision1 of this fixture is independently compared above.']
result['database'] = {k: db_after[k] for k in ['database', 'projects', 'forms', 'migration_count']}

assignment = json.loads((out / 'recipient-scoped-assignment.json').read_text())
check('recipient_active_internal_account', assignment['user_status'] == 'active' and not assignment['user_deleted'] and all(s['status'] == 'active' for s in assignment['staff']))
check('recipient_pm_only_in_test_project', len(assignment['roles']) == 1 and assignment['roles'][0]['role'] == 'pm' and assignment['roles'][0]['project_id'] == state['project_id'])
check('recipient_active_membership_only_in_test_project', len(assignment['memberships']) == 1 and assignment['memberships'][0]['active'] and assignment['memberships'][0]['project_id'] == state['project_id'])

manifest = json.loads((out / 'source-manifest.json').read_text())
mismatches = [p for p, h in manifest.items() if hashlib.sha256((root / p).read_bytes()).hexdigest() != h]
check('application_source_manifest_no_drift', not mismatches)
result['source_file_count'] = len(manifest)
result['head'] = subprocess.check_output(['git', 'rev-parse', 'HEAD'], cwd=root, text=True).strip()
result['files'] = {}
books = {}
archives = {}
for download in json.loads((out / 'browser-downloads.json').read_text()):
    path = out / download['file']
    raw = path.read_bytes()
    check(download['file'] + '_original_download_byte_equal', raw == Path(download['original']).read_bytes())
    book = load_workbook(path)
    books[download['file']] = book
    with ZipFile(path) as archive:
        archives[download['file']] = {n: archive.read(n) for n in archive.namelist()}
    errors = [(sheet.title, cell.coordinate, cell.value) for sheet in book for row in sheet for cell in row if cell.data_type == 'e']
    check(download['file'] + '_no_error_cells', not errors)
    media = []
    for name, data in archives[download['file']].items():
        if not name.startswith('xl/media/'):
            continue
        image = Image.open(io.BytesIO(data))
        image.load()
        media.append({'path': name, 'sha256': hashlib.sha256(data).hexdigest(), 'width': image.width, 'height': image.height, 'pixel_sha256': hashlib.sha256(image.convert('RGB').tobytes()).hexdigest()})
    check(download['file'] + '_one_decodable_embedded_photo', len(media) == 1)
    result['files'][download['file']] = {'sha256': hashlib.sha256(raw).hexdigest(), 'bytes': len(raw), 'sheets': book.sheetnames, 'media': media, 'progress': book[book.sheetnames[0]]['D10'].value, 'field_notes': book[book.sheetnames[0]]['A38'].value, 'caption': book['Hình ảnh']['A1'].value}

old_before = archives['BM02-lan-1-truoc.xlsx']
old_after = archives['BM02-lan-1-sau.xlsx']
different_parts = [p for p in sorted(set(old_before) | set(old_after)) if old_before.get(p) != old_after.get(p)]
result['revision1_archive_differences'] = different_parts
check('revision1_all_workbook_parts_equal_except_content_type_declaration_order', set(different_parts) <= {'[Content_Types].xml'})
# Excelize writes MIME declarations from a Go map. XML declaration order has
# no workbook meaning; compare every declaration and reject duplicates.
def content_types(raw):
    declarations = list(ET.fromstring(raw))
    normalized = sorted((el.tag, tuple(sorted(el.attrib.items())), el.text or '', el.tail or '') for el in declarations)
    assert len(normalized) == len(set(normalized)), 'duplicate content type declaration'
    return normalized
check('revision1_content_type_declarations_equal', content_types(old_before['[Content_Types].xml']) == content_types(old_after['[Content_Types].xml']))
for key in ['progress', 'field_notes', 'caption', 'media', 'sheets']:
    check('revision1_' + key + '_unchanged', result['files']['BM02-lan-1-truoc.xlsx'][key] == result['files']['BM02-lan-1-sau.xlsx'][key])
check('revision1_correct_progress', result['files']['BM02-lan-1-sau.xlsx']['progress'] == '25%')
check('revision2_correct_progress', result['files']['BM02-lan-2.xlsx']['progress'] == '60%')
check('revision2_distinct_embedded_photo', result['files']['BM02-lan-1-sau.xlsx']['media'][0]['pixel_sha256'] != result['files']['BM02-lan-2.xlsx']['media'][0]['pixel_sha256'])

recipient_token = json.loads(api.PRIVATE.read_text())['recipient']['token']
result['protected_photos'] = []
for rev, evidence_id in [(1, state['photo1_id']), (2, state['photo2_id'])]:
    req = urllib.request.Request(api.BASE + f'/work-reports/{state["work_report_id"]}/revisions/{rev}/evidence/{evidence_id}/file', headers={'Authorization': 'Bearer ' + recipient_token})
    with urllib.request.urlopen(req, timeout=30) as response:
        raw = response.read()
    original = (out / f'BM02-Lan-{rev}.png').read_bytes()
    check(f'recipient_protected_original_revision{rev}_matches_upload', raw == original)
    result['protected_photos'].append({'revision': rev, 'evidence_id': evidence_id, 'sha256': hashlib.sha256(raw).hexdigest()})

result['google'] = []
for number, name in [(1, 'BM02-lan-1-sau.xlsx'), (2, 'BM02-lan-2.xlsx')]:
    google = json.loads((out / f'google-{number}.json').read_text())['cells']['value']['structuredContent']
    cells = {}
    errors = []
    for sheet in google['sheets']:
        sheet_name = sheet['properties']['title']
        cells[sheet_name] = {}
        for block in sheet.get('data', []):
            for row_no, row in enumerate(block.get('rowData', []), block.get('startRow', 0) + 1):
                for col_no, cell in enumerate(row.get('values', []), block.get('startColumn', 0) + 1):
                    from openpyxl.utils import get_column_letter
                    address = get_column_letter(col_no) + str(row_no)
                    cells[sheet_name][address] = cell
                    if cell.get('effectiveValue', {}).get('errorValue'):
                        errors.append((sheet_name, address, cell['effectiveValue']))
    check(f'google_revision{number}_no_effective_errors', not errors)
    book = books[name]
    differences = []
    populated = 0
    for sheet in book:
        for row in sheet:
            for cell in row:
                if cell.value is None or cell.value == '':
                    continue
                populated += 1
                google_cell = cells[sheet.title].get(cell.coordinate, {})
                value = google_cell.get('userEnteredValue', {})
                expected = value.get('formulaValue') if cell.data_type == 'f' else value.get('numberValue') if isinstance(cell.value, (int, float)) else value.get('stringValue')
                if expected != cell.value:
                    differences.append((sheet.title, cell.coordinate, cell.value, expected))
    check(f'google_revision{number}_all_populated_cells_and_formulas_match_xlsx', not differences)
    total = cells['2026-10-12_01']['C28'].get('effectiveValue', {}).get('numberValue')
    check(f'google_revision{number}_crew_sum_recalculated', total == (3 if number == 1 else 5))
    result['google'].append({'revision': number, 'url': google['spreadsheetUrl'], 'ranges': ["'2026-10-12_01'!A1:H50", "'Hình ảnh'!A1:H5"], 'populated_cells_compared': populated, 'effective_error_count': len(errors), 'crew_sum': total})

api.save('acceptance-checks.json', result)
print(json.dumps({'checks_passed': len(result['checks']), 'archive_differences': different_parts, 'database': result['database'], 'source_file_count': len(manifest), 'google': result['google']}, ensure_ascii=False, indent=2))
