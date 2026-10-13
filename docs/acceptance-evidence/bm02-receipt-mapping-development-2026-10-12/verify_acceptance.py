"""Read canonical Development records and unchanged UI-exported workbooks.

Writes evidence JSON, not XLSX, Sheet content or application records. Requires
bundled Python/Pillow/openpyxl and the private temporary canonical API helper.
"""
import hashlib
from decimal import Decimal
import importlib.util
import io
import json
from pathlib import Path
import subprocess
from zipfile import ZipFile
import xml.etree.ElementTree as ET

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from PIL import Image

spec = importlib.util.spec_from_file_location('uat', '/private/tmp/consflow-bm02-receipt-mapping-uat.py')
uat = importlib.util.module_from_spec(spec)
spec.loader.exec_module(uat)
out, api, fixture, state = uat.out, uat.api, uat.fixture, uat.state
result = {'scope': 'Development BM02 project/date receipt mapping; existing canonical fixture and new form/delivery revision3', 'checks': {}}


def read(name):
    return json.loads((out / name).read_text())


def check(name, value):
    result['checks'][name] = bool(value)
    assert value, name


uat.renew()
form, report, receipt = uat.snapshot('final')
baseline_form, baseline_report = read('baseline-form-after.json'), read('baseline-delivery-after.json')
old_revisions = {row['revision_no']: row for row in baseline_report['revisions']}
current_revisions = {row['revision_no']: row for row in report['revisions']}
check('same_form_project_owner_submitted_revision3', form['instance']['id'] == state['form_id'] and form['instance']['project_id'] == state['project_id'] and form['instance']['work_item_id'] == state['report_work_item_id'] and form['instance']['status'] == 'submitted' and form['instance']['current_revision'] == 3)
check('same_report_delivery_revision3', report['work_report_id'] == state['work_report_id'] and report['current_revision_no'] == 3 and len(report['revisions']) == 3)
check('old_delivery_entire_records_unchanged', all(current_revisions[number] == row for number, row in old_revisions.items()))
new_events = {row['id']: row for row in form['events']}
check('all_earlier_form_events_unchanged', all(new_events.get(row['id']) == row for row in baseline_form['events']))
check('same_existing_field_definition', form['fields'] == baseline_form['fields'] and form['instance']['form_version_id'] == baseline_form['instance']['form_version_id'])
check('no_manual_arrivals_or_new_answer_keys', not read('form3-amend-request.json')['answers'].get('material_arrival_items') and set(form['answers']) <= set(baseline_form['answers']))
check('recipient_valid_requested_changes_revision2', any(row['event_type'] == 'needs_changes' and row['actor_id'] == state['recipient_id'] for row in report['events']))
check('receipt_unchanged_from_baseline', receipt == read('baseline-receipt-after.json'))
check('receipt_report_distinct_canonical_owners', receipt['work_item_id'] == state['receipt_work_item_id'] != form['instance']['work_item_id'])
facts = next(row['metadata_json']['native_daily_sources'] for row in form['events'] if row['event_type'] == 'amended' and row.get('metadata_json', {}).get('native_daily_sources', {}).get('form_revision_no') == 3)
arrivals = facts['structured_json']['material_arrival_items']
check('form3_one_correct_delivered_receipt12', len(arrivals) == 1 and arrivals[0]['source_id'] == state['receipt_id'] and arrivals[0]['source_item_id'] == receipt['items'][0]['id'] and arrivals[0]['source_work_item_id'] == state['receipt_work_item_id'] and Decimal(str(arrivals[0]['quantity'])) == Decimal('12') and arrivals[0]['unit'] == 'm2')
check('delivery3_exact_form_revision_arrivals', current_revisions[3]['snapshot_json']['structured_json']['material_arrival_items'] == arrivals)
check('old_deliveries_not_backfilled', all(not row['snapshot_json']['structured_json'].get('material_arrival_items') for row in old_revisions.values()))
check('new_delivery_exact_existing_photo2', [row['id'] for row in current_revisions[3]['evidence_records']] == [state['photo2_id']])

db_before, db_after = read('database-before.json'), fixture.db_snapshot()
uat.save('database-after.json', db_after)
allowed_changed = {'work_items': {state['report_work_item_id']}, 'workflow_form_instances': {state['form_id']}, 'operation_daily_reports': {state['source_id']}}
for table in db_before:
    if table in ['database', 'migration_count']:
        continue
    after_by_id = {row['id']: row for row in db_after[table]}
    check('preserved_prior_' + table, all(after_by_id.get(row['id']) == row for row in db_before[table] if row['id'] not in allowed_changed.get(table, set())))
check('database_migrations_and_record_counts', db_after['database'] == db_before['database'] == 'consflow' and db_after['migration_count'] == db_before['migration_count'] == 157 and len(db_after['projects']) == len(db_before['projects']) == 6 and len(db_after['workflow_form_instances']) == len(db_before['workflow_form_instances']) == 21)
result['database'] = {'name': db_after['database'], 'migration_count': db_after['migration_count'], 'projects': len(db_after['projects']), 'forms': len(db_after['workflow_form_instances']), 'allowed_changed_ids': {table: sorted(ids) for table, ids in allowed_changed.items()}}
manifest = read('source-manifest.json')
check('tested_built_source_manifest_unchanged', all(hashlib.sha256((fixture.ROOT / name).read_bytes()).hexdigest() == digest for name, digest in manifest.items()))
result['source_file_count'] = len(manifest)
result['head'] = subprocess.check_output(['git', 'rev-parse', 'HEAD'], cwd=fixture.ROOT, text=True).strip()
result['files'], books, archives = {}, {}, {}
for download in read('browser-downloads.json'):
    phase = download['phase']
    path = out / download['file']
    raw = path.read_bytes()
    check(phase + '_original_browser_bytes', raw == Path(download['original']).read_bytes())
    book = load_workbook(path)
    books[phase] = book
    with ZipFile(path) as archive:
        archives[phase] = {name: archive.read(name) for name in archive.namelist()}
    check(phase + '_two_sheets_no_stored_errors', book.sheetnames == ['2026-10-12_01', 'Hình ảnh'] and not [cell for sheet in book for row in sheet for cell in row if cell.data_type == 'e'])
    media = []
    for part, value in archives[phase].items():
        if part.startswith('xl/media/'):
            image = Image.open(io.BytesIO(value)).convert('RGB')
            media.append({'part': part, 'sha256': hashlib.sha256(value).hexdigest(), 'pixel_sha256': hashlib.sha256(image.tobytes()).hexdigest(), 'width': image.width, 'height': image.height})
    original_phase = 1 if phase == 'delivery1' else 2
    original = Image.open(fixture.OUT / f'BM02-workflow-Lan-{original_phase}.png').convert('RGB')
    check(phase + '_one_exact_original_photo', len(media) == 1 and media[0]['pixel_sha256'] == hashlib.sha256(original.tobytes()).hexdigest())
    sheet = book['2026-10-12_01']
    is_new = phase in ['form3', 'delivery3']
    if is_new:
        check(phase + '_material_name_quantity12', sheet['E26'].value == receipt['items'][0]['material_name'] and sheet['G26'].value == arrivals[0]['quantity'] and Decimal(str(sheet['G26'].value)) == Decimal('12'))
        check(phase + '_material_criterion_has_real_facts', receipt['items'][0]['material_name'] in sheet['D12'].value and '12' in sheet['D12'].value and 'm2' in sheet['D12'].value)
        check(phase + '_supplier_not_invented_origin', sheet['F26'].value in [None, ''])
    else:
        check(phase + '_old_receipt_missing_preserved', sheet['D12'].value == 'Chưa có ghi nhận vật tư nhập trường trong nguồn báo cáo.' and sheet['E25'].value in [None, ''] and sheet['G25'].value in [None, ''])
    check(phase + '_no_invented_progress_assessment', sheet['D10'].value == 'Chưa có ghi nhận tiến độ trong nguồn báo cáo.' and all(sheet[f'C{row}'].value == 'Chưa đánh giá' for row in range(10, 22)))
    result['files'][phase] = {'file': download['file'], 'sha256': hashlib.sha256(raw).hexdigest(), 'bytes': len(raw), 'media': media, 'material_note': sheet['D12'].value, 'crew': sheet['C25'].value}


def mime(raw):
    declarations = sorted((el.tag, tuple(sorted(el.attrib.items()))) for el in ET.fromstring(raw))
    assert len(declarations) == len(set(declarations))
    return declarations


def xml_value(element):
    return (element.tag, tuple(sorted(element.attrib.items())), element.text or '', tuple(xml_value(child) for child in element))


def resolved_sheet(raw, shared_strings):
    # Excelize's Go-map iteration can reorder the shared-string table. Resolve
    # every referenced index before comparing the full worksheet structure.
    # Styles, values, formulas, geometry and drawing references still must match.
    tree = ET.fromstring(raw)
    namespace = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
    strings = list(ET.fromstring(shared_strings))
    for cell in tree.iter(namespace + 'c'):
        if cell.get('t') == 's':
            value = cell.find(namespace + 'v')
            assert value is not None
            index = int(value.text)
            assert 0 <= index < len(strings)
            cell.remove(value)
            cell.set('t', 'resolvedSharedString')
            cell.append(strings[index])
    return xml_value(tree)


def archive_equal(before, after):
    if set(before) != set(after):
        return False
    for name in before:
        if before[name] == after[name]:
            continue
        if name == '[Content_Types].xml':
            if mime(before[name]) != mime(after[name]):
                return False
        elif name == 'xl/sharedStrings.xml':
            a, b = ET.fromstring(before[name]), ET.fromstring(after[name])
            if a.attrib != b.attrib or sorted(xml_value(child) for child in a) != sorted(xml_value(child) for child in b):
                return False
        elif name.startswith('xl/worksheets/') and name.endswith('.xml'):
            if resolved_sheet(before[name], before['xl/sharedStrings.xml']) != resolved_sheet(after[name], after['xl/sharedStrings.xml']):
                return False
        else:
            return False
    return True


result['old_archive_differences'] = {}
for phase, old_file in [('delivery1', 'BM02-lan-gui-1-sau.xlsx'), ('delivery2', 'BM02-lan-gui-2.xlsx')]:
    with ZipFile(fixture.OUT / old_file) as archive:
        baseline = {name: archive.read(name) for name in archive.namelist()}
    current = archives[phase]
    differences = [name for name in sorted(set(baseline) | set(current)) if baseline.get(name) != current.get(name)]
    check(phase + '_all_prior_parts_preserved_with_resolved_string_indices', archive_equal(baseline, current))
    result['old_archive_differences'][phase] = differences
check('delivery2_before_after_same_content_style_media', archive_equal(archives['delivery2-before'], archives['delivery2']))
check('form3_delivery3_template_cells_match', all(books['form3']['2026-10-12_01'].cell(row, col).value == books['delivery3']['2026-10-12_01'].cell(row, col).value for row in range(1, 29) for col in range(1, 9)))

result['google'] = []
for phase in ['form3', 'delivery3', 'delivery1', 'delivery2']:
    google = read('google-' + phase + '.json')['cells']
    observed, errors = {}, []
    for sheet in google['sheets']:
        title = sheet['properties']['title']
        observed[title] = {}
        for block in sheet.get('data', []):
            for row_no, row in enumerate(block.get('rowData', []), block.get('startRow', 0) + 1):
                for col_no, cell in enumerate(row.get('values', []), block.get('startColumn', 0) + 1):
                    address = get_column_letter(col_no) + str(row_no)
                    observed[title][address] = cell
                    if cell.get('effectiveValue', {}).get('errorValue'):
                        errors.append((title, address, cell['effectiveValue']))
    check(phase + '_google_no_effective_errors', not errors)
    differences, populated = [], 0
    for sheet in books[phase]:
        for row in range(1, 76 if sheet.title == '2026-10-12_01' else 6):
            for col in range(1, 9):
                cell = sheet.cell(row, col)
                expected = cell.value if cell.value not in [None, ''] else None
                value = observed[sheet.title].get(cell.coordinate, {}).get('userEnteredValue', {})
                actual = value.get('formulaValue', value.get('numberValue', value.get('stringValue')))
                actual = actual if actual not in [None, ''] else None
                populated += expected is not None
                if actual != expected:
                    differences.append((sheet.title, cell.coordinate, expected, actual))
    check(phase + '_google_all_bounded_cells_formulas_match', not differences)
    crew = observed['2026-10-12_01']['C28'].get('effectiveValue', {}).get('numberValue')
    check(phase + '_google_crew_recalculated', crew == (3 if phase == 'delivery1' else 6))
    result['google'].append({'phase': phase, 'url': google['spreadsheetUrl'], 'populated_cells_compared': populated, 'effective_error_count': len(errors), 'crew_sum': crew})
result['criteria'] = {'automatic_legitimate_receipt_reuse': 'PASS', 'frozen_form_delivery_and_old_history': 'PASS', 'original_browser_file_google_cells': 'PASS', 'user_visual_review': 'PENDING'}
result['status'] = 'TECHNICALLY VERIFIED within synthetic Development scope; user review pending'
result['screenshots'] = {path.name: hashlib.sha256(path.read_bytes()).hexdigest() for path in sorted(out.glob('*.jpg'))}
uat.save('acceptance-checks.json', result)
print(json.dumps({'status': result['status'], 'checks_passed': len(result['checks']), 'criteria': result['criteria'], 'google': result['google']}, ensure_ascii=False, indent=2))
