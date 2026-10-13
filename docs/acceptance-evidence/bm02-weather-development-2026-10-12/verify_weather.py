"""Read-only business-record and actual-download acceptance evidence.

Only routine auth renewal and evidence-file writes; never edits XLSX, Sheets,
canonical business data or stored revision snapshots.
"""
import datetime
import copy
from decimal import Decimal
import hashlib
import importlib.util
import json
from pathlib import Path
import subprocess
import sys
from zipfile import ZipFile
import xml.etree.ElementTree as ET

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from pypdf import PdfReader

spec = importlib.util.spec_from_file_location('receipt_uat', '/private/tmp/consflow-bm02-receipt-mapping-uat.py')
uat = importlib.util.module_from_spec(spec)
spec.loader.exec_module(uat)
ROOT = uat.fixture.ROOT
OUT = ROOT / 'docs/acceptance-evidence/bm02-weather-development-2026-10-12'


def save(name, data):
    (OUT / name).write_text(json.dumps(data, ensure_ascii=False, indent=2) + '\n')


def records(prefix):
    uat.renew()
    for name, path in [
        ('form', '/work-items/' + uat.state['report_work_item_id'] + '/form'),
        ('report', '/work-reports/daily-reports/' + uat.state['source_id']),
        ('receipt', '/incoming-acceptances/' + uat.state['receipt_id']),
    ]:
        save(prefix + '-' + name + '.json', uat.api.call(path))
    tables = ['projects', 'work_items', 'workflow_form_instances', 'workflow_form_events',
              'operation_daily_reports', 'work_reports', 'work_report_revisions', 'work_report_events',
              'incoming_acceptance_records', 'incoming_acceptance_items', 'evidence_records',
              'workflow_form_versions', 'workflow_form_fields', 'schema_migrations']
    fields = ','.join("'" + t + "',(SELECT COALESCE(json_agg(json_build_object('id',id,'row_md5',md5(row_to_json(r)::text)) ORDER BY id),'[]') FROM " + t + " r)" for t in tables if t != 'schema_migrations')
    sql = "SELECT json_build_object('database',current_database(),'migration_count',(SELECT count(*) FROM schema_migrations)," + fields + ");"
    raw = subprocess.check_output(['docker', 'exec', 'ttcons-postgres-1', 'sh', '-c', 'psql -U "$POSTGRES_USER" -d "$POSTGRES_DB" -At -c "$1"', 'sh', sql], text=True)
    save(prefix + '-database.json', json.loads(raw))
    print(json.dumps({'prefix': prefix, 'captured_at': datetime.datetime.now(datetime.timezone.utc).isoformat(), 'tables_hashed': len(tables) - 1}))


def manifest():
    previous = json.loads((ROOT / 'docs/acceptance-evidence/bm02-receipt-mapping-development-2026-10-12/source-manifest.json').read_text())
    # Extend the same application-source selection used by accepted receipt
    # evidence. A broader rg selection would mislabel existing seed/assets as new.
    files = set(previous) | {'backend/internal/core/construction_weather_presentation.go',
                             'backend/internal/core/construction_weather_presentation_test.go',
                             'frontend/src/modules/operations/constructionWeatherPresentation.ts',
                             'frontend/tests/constructionWeatherPresentation.test.ts'}
    values = {name: hashlib.sha256((ROOT / name).read_bytes()).hexdigest() for name in sorted(files)}
    save('source-manifest.json', values)
    save('source-baseline-manifest.json', previous)
    save('source-state.json', {'captured_at': datetime.datetime.now(datetime.timezone.utc).isoformat(), 'head': subprocess.check_output(['git','rev-parse','HEAD'], cwd=ROOT,text=True).strip(), 'file_count':len(values), 'manifest_sha256':hashlib.sha256((OUT/'source-manifest.json').read_bytes()).hexdigest(), 'changed_existing_files':[n for n in previous if n in values and values[n] != previous[n]], 'new_files':sorted(set(values)-set(previous)), 'missing_files':sorted(set(previous)-set(values))})


def verify():
    records('after')
    read = lambda name: json.loads((OUT / name).read_text())
    result = {'scope': 'Development weather presentation; existing synthetic BM02 fixture and immutable revisions', 'checks': {}, 'files': {}, 'google': []}

    def check(name, condition):
        result['checks'][name] = bool(condition)
        assert condition, name

    for name in ['form', 'report', 'receipt', 'database']:
        check(name + '_entire_baseline_preserved', read('before-' + name + '.json') == read('after-' + name + '.json'))
    form, report = read('after-form.json'), read('after-report.json')
    check('same_field_definition_11_fields', len(form['fields']) == 11)
    check('canonical_weather_codes_retained', form['answers']['weather'] == 'rain' and {r['revision_no']: r['snapshot_json']['weather'] for r in report['revisions']} == {1: 'hot', 2: 'rain', 3: 'rain'})
    check('no_false_resend_eligible_from_display_change', not report.get('delivery_source_changed'))
    source = read('source-manifest.json')
    check('built_source967_unchanged', len(source) == 967 and all(hashlib.sha256((ROOT / name).read_bytes()).hexdigest() == digest for name, digest in source.items()))
    result['source'] = read('source-state.json')
    result['database'] = {'name': read('after-database.json')['database'], 'migration_count': read('after-database.json')['migration_count'], 'projects': len(read('after-database.json')['projects']), 'forms': len(read('after-database.json')['workflow_form_instances']), 'business_tables_equal': 13}
    old_dir = ROOT / 'docs/acceptance-evidence/bm02-receipt-mapping-development-2026-10-12'
    old_files = {'form': 'BM02-phieu-phien-ban-3.xlsx', 'delivery3': 'BM02-lan-gui-3.xlsx', 'delivery1': 'BM02-lan-gui-1-giu-nguyen.xlsx', 'delivery2': 'BM02-lan-gui-2-giu-nguyen.xlsx'}
    books = {}
    n = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'

    def xml_tuple(el):
        return (el.tag, tuple(sorted(el.attrib.items())), el.text or '', el.tail or '', tuple(xml_tuple(child) for child in el))

    def resolved_xml(raw, strings):
        tree, pool = ET.fromstring(raw), list(ET.fromstring(strings))
        for cell in tree.iter(n + 'c'):
            if cell.get('t') == 's':
                value = cell.find(n + 'v')
                assert value is not None
                index = int(value.text)
                assert 0 <= index < len(pool)
                cell.remove(value)
                cell.set('t', 'resolvedSharedString')
                cell.append(copy.deepcopy(pool[index]))
            if cell.get('r') == 'B6':
                # Label correctness is checked independently against B6. Only
                # compare structure/style here; all other cells remain exact.
                for value in cell.iter(n + 't'):
                    value.text = '<weather-presentation>'
        return xml_tuple(tree)

    for download in read('browser-downloads.json'):
        phase, path = download['phase'], OUT / download['file']
        check(phase + '_actual_browser_bytes', path.read_bytes() == Path(download['original']).read_bytes())
        result['files'][phase] = {'file': download['file'], 'sha256': hashlib.sha256(path.read_bytes()).hexdigest(), 'bytes': path.stat().st_size}
        if path.suffix == '.pdf':
            pdf = PdfReader(path)
            text = '\n'.join(page.extract_text() for page in pdf.pages)
            label = 'Nắng nóng' if phase == 'pdf1' else 'Mưa'
            check(phase + '_actual_pdf_vietnamese_weather', label in text and len(pdf.pages) == 1)
            result['files'][phase].update(weather=label, pages=len(pdf.pages), progress_limitation='preexisting0% for missing workflow progress')
            continue
        book = load_workbook(path)
        books[phase] = book
        sheet = book.worksheets[0]
        check(phase + '_vietnamese_weather_B6', sheet['B6'].value == ('Nắng nóng' if phase == 'delivery1' else 'Mưa'))
        check(phase + '_no_stored_cell_errors', not [c.coordinate for s in book for row in s for c in row if c.data_type == 'e'])
        check(phase + '_progress_still_unknown', sheet['D10'].value == 'Chưa có ghi nhận tiến độ trong nguồn báo cáo.')
        if phase not in ['delivery1', 'delivery2']:
            check(phase + '_accepted_material12_preserved', sheet['G26'].value == '12.000' and Decimal(sheet['G26'].value) == Decimal('12'))
        if phase in old_files:
            old_path = old_dir / old_files[phase]
            old = load_workbook(old_path)
            check(phase + '_all_cells_styles_except_weather_equal', book.sheetnames == old.sheetnames and all(
                (book[s.title].cell(r,c).value == s.cell(r,c).value or (s.title == '2026-10-12_01' and r == 6 and c == 2))
                and book[s.title].cell(r,c)._style == s.cell(r,c)._style
                for s in old for r in range(1,max(s.max_row,book[s.title].max_row)+1) for c in range(1,max(s.max_column,book[s.title].max_column)+1)))
            with ZipFile(old_path) as a, ZipFile(path) as b:
                before, after = {k:a.read(k) for k in a.namelist()}, {k:b.read(k) for k in b.namelist()}
            check(phase + '_same_archive_parts', set(before) == set(after))
            mime_before, mime_after = ET.fromstring(before['[Content_Types].xml']), ET.fromstring(after['[Content_Types].xml'])
            declarations = lambda tree: sorted((c.tag,tuple(sorted(c.attrib.items()))) for c in tree)
            check(phase + '_mime_declarations_unchanged', mime_before.attrib == mime_after.attrib and declarations(mime_before) == declarations(mime_after))
            strings_before, strings_after = ET.fromstring(before['xl/sharedStrings.xml']), ET.fromstring(after['xl/sharedStrings.xml'])
            for entry in strings_before:
                if ''.join(t.text or '' for t in entry.iter(n+'t')) == old.worksheets[0]['B6'].value:
                    texts = list(entry.iter(n+'t'))
                    assert len(texts) == 1
                    texts[0].text = sheet['B6'].value
            check(phase + '_shared_strings_only_weather_changed', strings_before.attrib == strings_after.attrib and sorted(xml_tuple(c) for c in strings_before) == sorted(xml_tuple(c) for c in strings_after))
            check(phase + '_every_other_archive_part_equal', all(
                before[k] == after[k] for k in before if k != '[Content_Types].xml' and k != 'xl/sharedStrings.xml' and not (k.startswith('xl/worksheets/') and k.endswith('.xml'))))
            check(phase + '_worksheet_geometry_values_formulas_except_weather_equal', all(
                resolved_xml(before[k],before['xl/sharedStrings.xml']) == resolved_xml(after[k],after['xl/sharedStrings.xml'])
                for k in before if k.startswith('xl/worksheets/') and k.endswith('.xml')))
            result['files'][phase]['changed_cells_from_prior_acceptance'] = ['2026-10-12_01!B6']
        result['files'][phase]['weather'] = sheet['B6'].value

    for phase in ['form', 'delivery3', 'delivery1', 'delivery2']:
        google = read('google-' + phase + '.json')
        observed, errors = {}, []
        for sheet in google['cells']['sheets']:
            title = sheet['properties']['title']
            observed[title] = {}
            for block in sheet.get('data', []):
                for r, row in enumerate(block.get('rowData', []),block.get('startRow',0)+1):
                    for c, cell in enumerate(row.get('values', []),block.get('startColumn',0)+1):
                        observed[title][get_column_letter(c)+str(r)] = cell
                        if cell.get('effectiveValue',{}).get('errorValue'):
                            errors.append((title,r,c))
        check(phase + '_google_no_effective_error', not errors)
        differences, populated = [], 0
        for sheet in books[phase]:
            for r in range(1,76 if sheet.title == '2026-10-12_01' else 6):
                for c in range(1,9):
                    cell=sheet.cell(r,c)
                    expected=cell.value if cell.value not in [None,''] else None
                    value=observed[sheet.title].get(cell.coordinate,{}).get('userEnteredValue',{})
                    actual=value.get('formulaValue',value.get('numberValue',value.get('stringValue')))
                    actual=actual if actual not in [None,''] else None
                    populated += expected is not None
                    if actual != expected:
                        differences.append((sheet.title,cell.coordinate,expected,actual))
        check(phase + '_google_all_bounded_cells_formulas_equal', not differences)
        crew=observed['2026-10-12_01']['C28']['effectiveValue']['numberValue']
        check(phase + '_google_crew_recalculates', crew == (3 if phase == 'delivery1' else 6))
        result['google'].append({'phase':phase,'url':google['metadata']['spreadsheetUrl'],'populated_cells':populated,'errors':len(errors),'crew_sum':crew,'weather':observed['2026-10-12_01']['B6']['formattedValue']})
    result['images'] = {p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in sorted(OUT.iterdir()) if p.suffix in ['.jpg','.png']}
    result['criteria'] = {'weather_UI_XLSX_PDF': 'TECHNICALLY VERIFIED; user review pending', 'stored_history_no_new_field': 'PASS', 'prior_material_mapping_preserved': 'PASS', 'PDF_progress': 'separate preexisting0% issue, unapproved follow-up'}
    save('acceptance-checks.json',result)
    print(json.dumps({'checks':len(result['checks']),'images':len(result['images']),'criteria':result['criteria'],'google':result['google']},ensure_ascii=False,indent=2))


if __name__ == '__main__':
    if sys.argv[1] == 'verify':
        verify()
    elif sys.argv[1] == 'manifest':
        manifest()
    else:
        records(sys.argv[1])
